"""Frozen experimental batches and bounded, auditable measurement imports.

PostgreSQL is the production backend. SQLite is available only by an explicit
backend selection for local verification. Import confirmation is transactional;
the uploaded bytes and each source row are retained independently of normalized
values. These records establish what was measured, not biological efficacy.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
from pathlib import Path
import re
import sqlite3
import tempfile
from contextlib import contextmanager
from datetime import datetime, timezone
from typing import Any
import uuid
from xml.etree import ElementTree
import zipfile

MAX_UPLOAD_BYTES = 8 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 32 * 1024 * 1024
MAX_IMPORT_ROWS = 10000
MAX_COLUMNS = 48
MAX_CELL_CHARACTERS = 16384
MAX_CONTROLS = 64
RESULT_FIELDS = (
    "candidate_id", "endpoint", "value", "unit", "relation", "concentration",
    "concentration_unit", "conditions", "biological_context", "replicate_id",
    "status", "assay_id", "source_id",
)
EXPORT_FIELDS = (
    "order", "role", "candidate_id", "target_id", "parent_id", "smiles",
    "name", "priority_score_v2", "selection_lane", "null_fields", "row_json",
)
XLSX_CELL_LIMIT = 32767
XLSX_VALUE_FILE = "xlsx-cell-values.json"
XLSX_EMBEDDED_VALUE_FILE = "skinpipe/" + XLSX_VALUE_FILE
SUCCESS_STATUSES = {"success", "measured", "measured_success"}
FAILURE_STATUSES = {"failed", "unmeasurable", "not_measured"}
RELATIONS = {"=", "<", "<=", ">", ">="}
# No potency conversion is made across endpoints, mass units or assay contexts.
UNIT_MAP = {
    "M": ("nM", 1e9), "mM": ("nM", 1e6), "uM": ("nM", 1e3),
    "µM": ("nM", 1e3), "μM": ("nM", 1e3), "nM": ("nM", 1.0),
    "pM": ("nM", 1e-3), "%": ("%", 1.0), "fraction": ("fraction", 1.0),
    "ratio": ("ratio", 1.0), "unitless": ("unitless", 1.0),
    "binary": ("binary", 1.0), "mg/mL": ("mg/mL", 1.0),
    "ug/mL": ("ug/mL", 1.0), "µg/mL": ("ug/mL", 1.0),
    "ng/mL": ("ng/mL", 1.0), "RFU": ("RFU", 1.0), "RLU": ("RLU", 1.0),
    "pIC50": ("pIC50", 1.0), "pEC50": ("pEC50", 1.0),
    "pKi": ("pKi", 1.0), "pKd": ("pKd", 1.0),
}


class ExperimentError(ValueError):
    """An explicit validation, integrity or frozen-record conflict."""


def canonical_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True,
                      separators=(",", ":"), allow_nan=False)


def _sha(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _text(value: Any, name: str, *, required: bool = True) -> str:
    if value is None and not required:
        return ""
    if not isinstance(value, str) or (required and not value.strip()):
        raise ExperimentError(f"{name} must be a nonempty string")
    if len(value) > MAX_CELL_CHARACTERS or any(ord(c) < 32 and c not in "\n\t\r" for c in value):
        raise ExperimentError(f"{name} contains unsupported control characters or is too long")
    return value


def _parent_key(smiles: str) -> str:
    from .identity import chemical_identity

    identity = chemical_identity(smiles)
    if identity["structure_status"] != "valid":
        raise ExperimentError("An experimental batch requires a valid molecular structure")
    if identity["parent_ambiguous"]:
        raise ExperimentError("Resolve multi-organic-component material identity before experimental batching")
    return identity["parent_id"].split(":", 1)[1]


def _batch_rows(target_id: str, candidates: list[dict], controls: list[dict]) -> tuple[list[dict], list[dict]]:
    if not isinstance(candidates, list) or not isinstance(controls, list):
        raise ExperimentError("candidates and controls must be lists")
    if len(candidates) > 96 or len(controls) > MAX_CONTROLS:
        raise ExperimentError("A batch permits at most 96 candidates and 64 separate controls")
    ids: set[str] = set()
    parents: set[str] = set()
    structures: set[str] = set()
    result: list[list[dict]] = [[], []]
    for position, rows in enumerate((candidates, controls)):
        for row in rows:
            if not isinstance(row, dict):
                raise ExperimentError("Every batch row must be an object")
            item = json.loads(canonical_json(row))
            cid = _text(item.get("candidate_id"), "candidate_id")
            if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.:-]{0,255}", cid):
                raise ExperimentError("candidate_id must be a stable alphanumeric machine identifier")
            parent = _text(item.get("parent_id"), "parent_id")
            if _text(item.get("target_id"), "target_id") != target_id:
                raise ExperimentError("Every candidate and control must belong to this target")
            if cid in ids or parent in parents:
                raise ExperimentError("Candidate IDs and parent IDs must be unique, including controls")
            key = _parent_key(_text(item.get("smiles"), "smiles"))
            if parent != "parent:" + key:
                raise ExperimentError("parent_id does not match the versioned structure identity")
            if key in structures:
                raise ExperimentError("Salt or supplier variants cannot inflate the unique parent count")
            ids.add(cid)
            parents.add(parent)
            structures.add(key)
            item["counting_parent_sha256"] = key
            result[position].append(item)
    return result[0], result[1]


def freeze_payload(payload: dict) -> dict:
    """Return a new content-bound snapshot without mutating the draft."""
    from rdkit import rdBase
    from .identity import IDENTITY_VERSION, PARENT_VERSION

    target = _text(payload.get("target_id"), "target_id")
    _text(payload.get("snapshot_id"), "snapshot_id")
    if not isinstance(payload.get("policy"), dict) or not payload["policy"]:
        raise ExperimentError("A nonempty frozen selection policy is required")
    candidates, controls = _batch_rows(target, payload["candidates"], payload.get("controls", []))
    if payload["policy"].get("procurement_requirement") == "dated-exact-form-v1":
        from .procurement import documented_offers
        held = [row["candidate_id"] for row in candidates if not documented_offers(row)]
        if held:
            raise ExperimentError("Final purchase list requires dated exact-form supplier evidence: " + ", ".join(held))
    if not candidates:
        raise ExperimentError("A frozen experimental batch requires at least one candidate")
    result = json.loads(canonical_json(payload))
    result.pop("payload_sha256", None)
    result.update(schema_version=1, status="frozen", candidates=candidates,
                  controls=controls, candidate_count=len(candidates), control_count=len(controls))
    if payload["policy"].get("procurement_requirement") == "dated-exact-form-v1":
        result["experiment_readiness"] = {"status": "open_items", "items": ["assay_conditions", "controls", "stock_and_lead_time"],
                                          "candidate_list_frozen": True, "trial_ready": False}
    result["identity_contract"] = {"identity_version": IDENTITY_VERSION,
                                   "parent_policy": PARENT_VERSION, "rdkit_version": rdBase.rdkitVersion}
    result["payload_sha256"] = _sha(canonical_json(result).encode())
    return result


def _verify_frozen(payload: dict) -> None:
    content = {k: v for k, v in payload.items() if k != "payload_sha256"}
    if payload.get("status") != "frozen" or payload.get("payload_sha256") != _sha(canonical_json(content).encode()):
        raise ExperimentError("Frozen batch content failed integrity verification")


def _spreadsheet_text(value: Any) -> Any:
    if isinstance(value, str) and value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _csv_text(value: Any) -> Any:
    """Display NUL as JSON's visible escape; row_json retains exact originals.

    Python 3.10's CSV reader/writer rejects NUL even inside quoted fields.
    Other values retain their existing formula-safe display byte for byte.
    """
    displayed = _spreadsheet_text(value)
    return displayed.replace("\x00", "\\u0000") if isinstance(displayed, str) else displayed


def _xlsx_value(value: Any, cell_key: str, values: dict) -> Any:
    """Keep a safe inline display or a short reference to the complete value."""
    displayed = _spreadsheet_text(value)
    if not isinstance(displayed, str):
        return displayed
    # Use UTF-16 units conservatively for Excel, and check after formula escaping
    # so a 32767-character formula label cannot lose its final character.
    xml_unsafe = any(not (ord(c) in (9, 10, 13) or 32 <= ord(c) <= 0xD7FF
                         or 0xE000 <= ord(c) <= 0xFFFD or 0x10000 <= ord(c) <= 0x10FFFF)
                     for c in displayed)
    reserved_reference = displayed.startswith('{"$skinpipe_xlsx_cell":')
    if not xml_unsafe and not reserved_reference and len(displayed.encode("utf-16-le")) // 2 <= XLSX_CELL_LIMIT:
        return displayed
    checksum = _sha(canonical_json(value).encode())
    values["cells"][cell_key] = {"value": value, "sha256": checksum}
    return canonical_json({"$skinpipe_xlsx_cell": {
        "file": XLSX_VALUE_FILE, "cell": cell_key, "sha256": checksum,
    }})


def restore_export_xlsx_value(value: Any, values: dict | None, *,
                              payload_sha256: str, spreadsheet_safe: bool = True) -> Any:
    """Resolve an export reference with source/value integrity checks.

    Ordinary cells retain their displayed values. Set spreadsheet_safe=False
    to recover an external original string; the default retains formula guards.
    The sidecar is included beside the workbook and inside its ZIP container.
    """
    if not isinstance(value, str) or not value.startswith('{"$skinpipe_xlsx_cell":'):
        return value
    try:
        marker = json.loads(value)
        reference = marker["$skinpipe_xlsx_cell"]
        if (set(marker) != {"$skinpipe_xlsx_cell"} or not isinstance(reference, dict)
                or set(reference) != {"file", "cell", "sha256"}
                or reference["file"] != XLSX_VALUE_FILE):
            raise ExperimentError("Invalid XLSX value reference")
        if (not isinstance(values, dict) or values.get("schema") != "skinpipe-xlsx-values-v1"
                or values.get("payload_sha256") != payload_sha256):
            raise ExperimentError("XLSX value reference source payload mismatch")
        item = values.get("cells", {}).get(reference["cell"])
        if not isinstance(item, dict) or "value" not in item:
            raise ExperimentError("XLSX referenced value is missing")
        original = item["value"]
        if (not isinstance(original, str) or _sha(canonical_json(original).encode()) != reference["sha256"]
                or item.get("sha256") != reference["sha256"]):
            raise ExperimentError("XLSX referenced value checksum mismatch")
        return _spreadsheet_text(original) if spreadsheet_safe else original
    except (json.JSONDecodeError, KeyError, TypeError, AttributeError) as error:
        raise ExperimentError("Invalid XLSX value reference") from error


def _export_rows(payload: dict) -> list[dict]:
    rows = []
    for role in ("candidates", "controls"):
        for index, row in enumerate(payload[role], 1):
            item = {key: row.get(key) for key in EXPORT_FIELDS}
            item.update(order=index, role="candidate" if role == "candidates" else "control",
                        selection_lane=row.get("selection_lane", row.get("selection_group")),
                        null_fields=canonical_json(sorted(k for k, v in row.items() if v is None)),
                        row_json=canonical_json(row))
            rows.append(item)
    return rows


def export_payload(payload: dict) -> dict[str, bytes]:
    """Produce ordered CSV/XLSX/SDF and a hash-manifest evidence ZIP in memory."""
    _verify_frozen(payload)
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from rdkit import Chem

    rows = _export_rows(payload)
    csv_buffer = io.StringIO(newline="")
    writer = csv.DictWriter(csv_buffer, fieldnames=EXPORT_FIELDS)
    writer.writeheader()
    writer.writerows({k: _csv_text(v) for k, v in row.items()} for row in rows)
    workbook = Workbook()
    xlsx_values = {"schema": "skinpipe-xlsx-values-v1",
                   "payload_sha256": payload["payload_sha256"], "cells": {}}
    sheet = workbook.active
    sheet.title = "Experiment batch"
    sheet.append(list(EXPORT_FIELDS))
    for row in rows:
        number = sheet.max_row + 1
        sheet.append([_xlsx_value(row[key], sheet.title + "!" + get_column_letter(column) + str(number), xlsx_values)
                      for column, key in enumerate(EXPORT_FIELDS, 1)])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    metadata = workbook.create_sheet("Frozen policy")
    for number, key in enumerate(("batch_id", "target_id", "snapshot_id", "payload_sha256", "policy"), 1):
        value = payload.get(key)
        value = canonical_json(value) if isinstance(value, dict) else value
        metadata.append([key, _xlsx_value(value, metadata.title + "!B" + str(number), xlsx_values)])
    sidecar = None
    if xlsx_values["cells"]:
        sidecar = canonical_json(xlsx_values).encode()
        instructions = workbook.create_sheet("XLSX value references")
        instructions.append(["schema", "skinpipe-xlsx-values-v1"])
        instructions.append(["source payload SHA256", payload["payload_sha256"]])
        instructions.append(["complete original values", XLSX_VALUE_FILE])
        instructions.append(["embedded copy in workbook ZIP", XLSX_EMBEDDED_VALUE_FILE])
        instructions.append(["sidecar SHA256", _sha(sidecar)])
        instructions.append(["reference format", "A $skinpipe_xlsx_cell JSON reference names a cell and SHA256 of its complete original JSON string."])
        instructions.append(["recovery", "Verify the source payload and value SHA256, restore the string, then parse row_json or policy JSON as needed."])
        instructions.append(["formula safety", "Original labels are preserved. When writing restored values to a spreadsheet, keep formula escaping enabled."])
        instructions.append(["CSV control characters", "NUL is displayed as the visible JSON escape \\u0000 for portable CSV readers; exact originals remain in row_json, batch.json and XLSX references."])
    xlsx_buffer = io.BytesIO()
    workbook.save(xlsx_buffer)
    workbook.close()
    xlsx_bytes = xlsx_buffer.getvalue()
    if sidecar is not None:
        embedded = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(xlsx_bytes)) as original, zipfile.ZipFile(embedded, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for entry in original.infolist():
                data = original.read(entry.filename)
                if entry.filename == "[Content_Types].xml":
                    types = ElementTree.fromstring(data)
                    ElementTree.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Default",
                                           {"Extension": "json", "ContentType": "application/json"})
                    data = ElementTree.tostring(types, encoding="utf-8", xml_declaration=True)
                archive.writestr(entry, data)
            archive.writestr(XLSX_EMBEDDED_VALUE_FILE, sidecar)
        xlsx_bytes = embedded.getvalue()
    # The Python StringIO adapter decodes each native write separately, which
    # can split a long UTF-8 value mid-character. A private native file stream
    # preserves the bytes; it is removed before this in-memory API returns.
    with tempfile.TemporaryDirectory(prefix="skinpipe-batch-sdf-") as temporary:
        sdf_path = Path(temporary) / "candidates.sdf"
        sdf_writer = Chem.SDWriter(str(sdf_path))
        try:
            for row in rows:
                mol = Chem.MolFromSmiles(row["smiles"])
                if mol is None:
                    raise ExperimentError("Frozen structure can no longer be parsed")
                mol.SetProp("_Name", row["candidate_id"])
                for key, value in row.items():
                    if value is not None:
                        mol.SetProp(key, str(value))
                sdf_writer.write(mol)
        finally:
            sdf_writer.close()
        sdf_bytes = sdf_path.read_bytes()
    artifacts = {
        "candidates.csv": csv_buffer.getvalue().encode("utf-8-sig"),
        "candidates.xlsx": xlsx_bytes,
        "candidates.sdf": sdf_bytes,
        "batch.json": canonical_json(payload).encode(),
    }
    if sidecar is not None:
        artifacts[XLSX_VALUE_FILE] = sidecar
    template = io.StringIO(newline="")
    template_writer = csv.DictWriter(template, fieldnames=RESULT_FIELDS)
    template_writer.writeheader()
    for row in rows:
        template_writer.writerow({"candidate_id": _spreadsheet_text(row["candidate_id"]), "status": "not_measured"})
    artifacts["results-template.csv"] = template.getvalue().encode("utf-8-sig")
    manifest = {name: {"sha256": _sha(data), "bytes": len(data)} for name, data in artifacts.items()}
    bundle = io.BytesIO()
    with zipfile.ZipFile(bundle, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in artifacts.items():
            archive.writestr(name, data)
        archive.writestr("manifest.json", canonical_json(manifest))
    artifacts["evidence.zip"] = bundle.getvalue()
    return artifacts


def parse_result_upload(raw: bytes, filename: str) -> list[dict]:
    """Read bounded CSV or one-sheet XLSX; reject formulas and external links."""
    if not isinstance(raw, bytes) or not raw or len(raw) > MAX_UPLOAD_BYTES:
        raise ExperimentError("Upload must contain 1 byte to 8 MiB")
    suffix = Path(filename).suffix.lower()
    source_rows: list[tuple[int, list[Any]]] = []
    sheet_name = None
    if suffix == ".csv":
        try:
            text = raw.decode("utf-8-sig")
            source_rows = []
            for index, cells in enumerate(csv.reader(io.StringIO(text, newline="")), 1):
                if index > MAX_IMPORT_ROWS + 1:
                    raise ExperimentError("Upload exceeds the row limit")
                source_rows.append((index, cells))
        except (UnicodeError, csv.Error) as exc:
            raise ExperimentError("CSV must be valid UTF-8 with parseable cells") from exc
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                files = archive.infolist()
                if len(files) > 256 or sum(f.file_size for f in files) > MAX_UNCOMPRESSED_BYTES:
                    raise ExperimentError("XLSX archive exceeds its decompression limit")
                if len({f.filename for f in files}) != len(files):
                    raise ExperimentError("XLSX has duplicate archive entries")
                for entry in files:
                    name = entry.filename.lower()
                    if entry.flag_bits & 1 or "externallinks/" in name or name.endswith("vbaproject.bin"):
                        raise ExperimentError("Encrypted content, external links and macros are not accepted")
                    if ".." in Path(entry.filename).parts or entry.filename.startswith("/"):
                        raise ExperimentError("Invalid archive member path")
                    if entry.file_size > max(1024 * 1024, entry.compress_size * 500):
                        raise ExperimentError("XLSX compression ratio exceeds its limit")
                    if name.endswith((".xml", ".rels")):
                        content = archive.read(entry)
                        if b"\x00" in content or b"<!DOCTYPE" in content.upper() or b"<!ENTITY" in content.upper():
                            raise ExperimentError("XML document types and entities are not accepted")
                        if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                            for _, element in ElementTree.iterparse(io.BytesIO(content), events=("start",)):
                                if element.tag.rsplit("}", 1)[-1] == "c":
                                    reference = re.fullmatch(r"([A-Z]+)([0-9]+)", element.get("r", ""))
                                    if reference is None:
                                        raise ExperimentError("Unsupported worksheet cell reference")
                                    column = 0
                                    for letter in reference[1]:
                                        column = column * 26 + ord(letter) - 64
                                    if column > MAX_COLUMNS or int(reference[2]) > MAX_IMPORT_ROWS + 1:
                                        raise ExperimentError("Worksheet exceeds the row or column limit")
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            try:
                if len(book.worksheets) != 1:
                    raise ExperimentError("Result imports require exactly one worksheet")
                sheet = book.worksheets[0]
                sheet_name = sheet.title
                # Declared dimensions are untrusted; cap iteration even when an
                # XLSX incorrectly claims a tiny or excessively large rectangle.
                sheet.reset_dimensions()
                for index, cells in enumerate(sheet.iter_rows(max_row=MAX_IMPORT_ROWS + 2, max_col=MAX_COLUMNS + 1), 1):
                    if index > MAX_IMPORT_ROWS + 1:
                        if any(c.value is not None for c in cells):
                            raise ExperimentError("Worksheet exceeds the row limit")
                        continue
                    if any(c.value is not None for c in cells[MAX_COLUMNS:]):
                        raise ExperimentError("Worksheet exceeds the column limit")
                    if any(c.data_type == "f" for c in cells):
                        raise ExperimentError(f"Formulas are not accepted at worksheet row {index}")
                    values = [c.value for c in cells]
                    while values and values[-1] is None:
                        values.pop()
                    if values:
                        source_rows.append((index, values))
            finally:
                book.close()
        except (zipfile.BadZipFile, KeyError, OSError, ValueError, ElementTree.ParseError) as exc:
            if isinstance(exc, ExperimentError):
                raise
            raise ExperimentError("Invalid or unsupported XLSX workbook") from exc
    else:
        raise ExperimentError("Only CSV and XLSX result files are accepted")
    if not source_rows:
        raise ExperimentError("Result file is empty")
    headers = [str(v).strip() if v is not None else "" for v in source_rows[0][1]]
    if len(headers) > MAX_COLUMNS or len(headers) != len(set(headers)) or not all(headers):
        raise ExperimentError("Column headers must be nonempty, unique and bounded")
    required = {"candidate_id", "endpoint", "value", "unit", "relation", "replicate_id", "status"}
    if not required.issubset(headers):
        raise ExperimentError("Required columns missing: " + ", ".join(sorted(required - set(headers))))
    result = []
    for number, cells in source_rows[1:]:
        if not any(v is not None and str(v).strip() for v in cells):
            continue
        if len(cells) > len(headers):
            raise ExperimentError(f"Unexpected extra cells at row {number}")
        values = ["" if value is None else str(value) for value in cells]
        if any(len(value) > MAX_CELL_CHARACTERS for value in values):
            raise ExperimentError(f"Cell exceeds size limit at row {number}")
        values.extend([""] * (len(headers) - len(values)))
        result.append({"source_row": {"row": number, "sheet": sheet_name},
                       "raw": dict(zip(headers, values))})
    if not result:
        raise ExperimentError("No nonempty result rows were supplied")
    return result


def _finite(value: str, name: str) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ExperimentError(f"{name} must be numeric") from exc
    if not math.isfinite(number):
        raise ExperimentError(f"{name} must be finite")
    return number


def validate_observations(rows: list[dict], payload: dict, *, existing_keys: set[str] | None = None) -> dict:
    _verify_frozen(payload)
    allowed = {row["candidate_id"] for row in payload["candidates"] + payload["controls"]}
    seen = set(existing_keys or ())
    normalized, errors, warnings = [], [], []
    for source in rows:
        raw = source["raw"]
        location = source["source_row"]
        try:
            row = {key: raw.get(key, "").strip() for key in RESULT_FIELDS}
            for key in ("candidate_id", "endpoint", "replicate_id", "status"):
                _text(row[key], key)
            if row["candidate_id"] not in allowed:
                raise ExperimentError("Unknown candidate ID for this frozen batch")
            status = row["status"].lower()
            if status not in SUCCESS_STATUSES | FAILURE_STATUSES:
                raise ExperimentError("Unsupported measurement status")
            relation = row["relation"] or ("=" if status in SUCCESS_STATUSES else "")
            if relation and relation not in RELATIONS:
                raise ExperimentError("Unsupported measurement relation")
            numeric = None
            unit = row["unit"] or None
            if status in SUCCESS_STATUSES:
                numeric = _finite(row["value"], "value")
                if unit not in UNIT_MAP:
                    raise ExperimentError("Unknown unit; supply an explicitly supported measurement unit")
                normalized_unit, factor = UNIT_MAP[unit]
                numeric *= factor
                if not math.isfinite(numeric):
                    raise ExperimentError("Normalized value exceeds the finite range")
                if normalized_unit == "binary" and numeric not in (0, 1):
                    raise ExperimentError("Binary observations require 0 or 1")
                if normalized_unit in {"nM", "mg/mL", "ug/mL", "ng/mL"} and numeric < 0:
                    raise ExperimentError("Concentration-valued measurements cannot be negative")
            else:
                if row["value"]:
                    raise ExperimentError("Failed or unmeasured observations must not contain a numeric result")
                normalized_unit = unit
            concentration = None
            concentration_unit = row["concentration_unit"] or None
            if row["concentration"]:
                concentration = _finite(row["concentration"], "concentration")
                if concentration < 0 or concentration_unit not in UNIT_MAP or UNIT_MAP[concentration_unit][0] not in {"nM", "mg/mL", "ug/mL", "ng/mL"}:
                    raise ExperimentError("Concentration requires a nonnegative value and concentration unit")
                concentration_unit, factor = UNIT_MAP[concentration_unit]
                concentration *= factor
                if not math.isfinite(concentration):
                    raise ExperimentError("Normalized concentration exceeds the finite range")
            elif concentration_unit:
                raise ExperimentError("Concentration unit was supplied without a concentration")
            item = {"target_id": payload["target_id"], "candidate_id": row["candidate_id"],
                    "endpoint": row["endpoint"], "value": numeric, "unit": normalized_unit,
                    "relation": relation or None, "concentration": concentration,
                    "concentration_unit": concentration_unit, "conditions": row["conditions"],
                    "biological_context": row["biological_context"], "replicate_id": row["replicate_id"],
                    "status": "measured" if status in SUCCESS_STATUSES else status,
                    "assay_id": row["assay_id"], "source_id": row["source_id"],
                    "kind": "measured" if status in SUCCESS_STATUSES else "observation_status"}
            key_fields = {key: item[key] for key in ("target_id", "candidate_id", "endpoint", "unit", "concentration",
                          "concentration_unit", "conditions", "biological_context", "replicate_id", "assay_id", "source_id")}
            key = _sha(canonical_json(key_fields).encode())
            if key in seen:
                raise ExperimentError("Duplicate observation key; use a distinct replicate or assay identifier")
            seen.add(key)
            if not item["assay_id"] or not item["biological_context"]:
                warnings.append({"source_row": location, "message": "Incomplete assay context; evidence remains context-limited"})
            normalized.append({"observation_key": key, "normalized": item, "raw": dict(raw), "source_row": location})
        except (ExperimentError, TypeError, ValueError) as exc:
            errors.append({"source_row": location, "message": str(exc)})
    return {"valid": not errors, "rows": normalized, "errors": errors, "warnings": warnings,
            "source_row_count": len(rows), "accepted_row_count": len(normalized)}


class ExperimentStore:
    def __init__(self, dsn: str, *, backend: str = "postgresql", artifact_root: str | Path):
        if backend not in {"postgresql", "sqlite"}:
            raise ExperimentError("Select postgresql or explicit local sqlite backend")
        if backend == "sqlite" and dsn == ":memory:":
            raise ExperimentError("SQLite verification requires a file to support separate transactions")
        self.dsn, self.backend = dsn, backend
        self.artifact_root = Path(artifact_root)
        self.artifact_root.mkdir(parents=True, exist_ok=True)
        with self._transaction() as connection:
            for statement in (
                "CREATE TABLE IF NOT EXISTS platform_experiment_batches (id TEXT PRIMARY KEY, target_id TEXT NOT NULL, status TEXT NOT NULL, payload TEXT NOT NULL, created_at TEXT NOT NULL)",
                "CREATE TABLE IF NOT EXISTS platform_experiment_imports (id TEXT PRIMARY KEY, batch_id TEXT NOT NULL REFERENCES platform_experiment_batches(id), raw_sha256 TEXT NOT NULL, filename TEXT NOT NULL, preview TEXT NOT NULL, created_at TEXT NOT NULL)",
                "CREATE TABLE IF NOT EXISTS platform_experiment_receipts (idempotency_key TEXT PRIMARY KEY, preview_id TEXT NOT NULL, batch_id TEXT NOT NULL, receipt TEXT NOT NULL)",
                "CREATE TABLE IF NOT EXISTS platform_experiment_observations (batch_id TEXT NOT NULL REFERENCES platform_experiment_batches(id), observation_key TEXT NOT NULL, candidate_id TEXT NOT NULL, preview_id TEXT NOT NULL, payload TEXT NOT NULL, PRIMARY KEY (batch_id, observation_key))",
            ):
                connection.execute(statement)

    @contextmanager
    def _transaction(self):
        if self.backend == "sqlite":
            connection = sqlite3.connect(self.dsn, timeout=30)
            connection.execute("PRAGMA foreign_keys=ON")
            connection.execute("BEGIN IMMEDIATE")
        else:
            try:
                import psycopg
            except ImportError as exc:
                raise RuntimeError("Production ExperimentStore requires psycopg; no SQLite fallback is performed") from exc
            connection = psycopg.connect(self.dsn)
        try:
            yield connection
            connection.commit()
        except BaseException:
            connection.rollback()
            raise
        finally:
            connection.close()

    def _sql(self, connection, statement: str, values: tuple = ()):
        return connection.execute(statement.replace("?", "%s") if self.backend == "postgresql" else statement, values)

    def _batch(self, connection, batch_id: str, *, lock: bool = False) -> dict:
        suffix = " FOR UPDATE" if lock and self.backend == "postgresql" else ""
        row = self._sql(connection, "SELECT payload FROM platform_experiment_batches WHERE id=?" + suffix, (batch_id,)).fetchone()
        if row is None:
            raise ExperimentError("Unknown experiment batch")
        payload = json.loads(row[0])
        if payload.get("status") == "frozen":
            _verify_frozen(payload)
        return payload

    def create_batch(self, target_id: str, candidates: list[dict], controls: list[dict], snapshot_id: str,
                     policy: dict, batch_id: str | None = None) -> dict:
        _text(target_id, "target_id")
        _text(snapshot_id, "snapshot_id")
        candidates, controls = _batch_rows(target_id, candidates, controls)
        payload = {"batch_id": batch_id or str(uuid.uuid4()), "target_id": target_id,
                   "snapshot_id": snapshot_id, "policy": policy, "candidates": candidates,
                   "controls": controls, "status": "draft", "created_at": _now(), "schema_version": 1}
        _text(payload["batch_id"], "batch_id")
        if len(payload["batch_id"]) > 256:
            raise ExperimentError("batch_id exceeds 256 characters")
        with self._transaction() as connection:
            self._sql(connection, "INSERT INTO platform_experiment_batches VALUES (?,?,?,?,?)",
                      (payload["batch_id"], target_id, "draft", canonical_json(payload), payload["created_at"]))
        return payload

    def get_batch(self, batch_id: str) -> dict:
        with self._transaction() as connection:
            return self._batch(connection, batch_id)

    def list_batches(self, target_id: str | None = None) -> list[dict]:
        with self._transaction() as connection:
            query = "SELECT payload FROM platform_experiment_batches"
            values = ()
            if target_id is not None:
                query += " WHERE target_id=?"
                values = (target_id,)
            rows = self._sql(connection, query + " ORDER BY created_at,id", values).fetchall()
            batches = []
            for row in rows:
                payload = json.loads(row[0])
                try:
                    if payload.get("status") == "frozen":
                        _verify_frozen(payload)
                    batches.append(payload)
                except ExperimentError:
                    batches.append({"batch_id": payload.get("batch_id"), "target_id": payload.get("target_id"),
                                    "status": "integrity_failed", "candidates": [], "controls": []})
            return batches

    def update_draft(self, batch_id: str, candidates: list[dict], controls: list[dict]) -> dict:
        with self._transaction() as connection:
            payload = self._batch(connection, batch_id, lock=True)
            if payload["status"] != "draft":
                raise ExperimentError("Frozen batches are immutable; create a new batch revision")
            payload["candidates"], payload["controls"] = _batch_rows(payload["target_id"], candidates, controls)
            self._sql(connection, "UPDATE platform_experiment_batches SET payload=? WHERE id=?",
                      (canonical_json(payload), batch_id))
            return payload

    def freeze_batch(self, batch_id: str) -> dict:
        with self._transaction() as connection:
            payload = self._batch(connection, batch_id, lock=True)
            if payload["status"] == "frozen":
                return payload
            payload["frozen_at"] = _now()
            frozen = freeze_payload(payload)
            self._sql(connection, "UPDATE platform_experiment_batches SET status='frozen',payload=? WHERE id=?",
                      (canonical_json(frozen), batch_id))
            return frozen

    def export_batch(self, batch_id: str) -> dict[str, bytes]:
        return export_payload(self.get_batch(batch_id))

    def _save_raw(self, raw: bytes) -> str:
        if not isinstance(raw, bytes) or not raw or len(raw) > MAX_UPLOAD_BYTES:
            raise ExperimentError("Upload must contain 1 byte to 8 MiB")
        digest = _sha(raw)
        path = self.artifact_root / digest
        temporary = self.artifact_root / (".upload-" + uuid.uuid4().hex)
        try:
            with temporary.open("xb") as stream:
                stream.write(raw)
                stream.flush()
                os.fsync(stream.fileno())
            try:
                os.link(temporary, path)
                directory = os.open(self.artifact_root, os.O_RDONLY)
                try:
                    os.fsync(directory)
                finally:
                    os.close(directory)
            except FileExistsError:
                if path.read_bytes() != raw:
                    raise ExperimentError("Content-addressed raw artifact failed integrity verification")
        finally:
            temporary.unlink(missing_ok=True)
        return digest

    def preview_import(self, batch_id: str, raw: bytes, filename: str) -> dict:
        payload = self.get_batch(batch_id)
        _verify_frozen(payload)
        digest = self._save_raw(raw)
        with self._transaction() as connection:
            existing = {r[0] for r in self._sql(connection, "SELECT observation_key FROM platform_experiment_observations WHERE batch_id=?", (batch_id,)).fetchall()}
        try:
            parsed = parse_result_upload(raw, filename)
            preview = validate_observations(parsed, payload, existing_keys=existing)
        except ExperimentError as exc:
            preview = {"valid": False, "rows": [], "errors": [{"source_row": None, "message": str(exc)}],
                       "warnings": [], "source_row_count": 0, "accepted_row_count": 0}
        preview.update(preview_id=str(uuid.uuid4()), batch_id=batch_id, batch_sha256=payload["payload_sha256"],
                       raw_sha256=digest, filename=Path(filename).name, created_at=_now())
        with self._transaction() as connection:
            self._sql(connection, "INSERT INTO platform_experiment_imports VALUES (?,?,?,?,?,?)",
                      (preview["preview_id"], batch_id, digest, Path(filename).name, canonical_json(preview), preview["created_at"]))
        return preview

    def get_import(self, preview_id: str) -> dict:
        with self._transaction() as connection:
            row = self._sql(connection, "SELECT preview FROM platform_experiment_imports WHERE id=?", (preview_id,)).fetchone()
            if row is None:
                raise ExperimentError("Unknown result import preview")
            return json.loads(row[0])

    def confirm_import(self, preview_id: str, idempotency_key: str) -> dict:
        _text(idempotency_key, "idempotency_key")
        if len(idempotency_key) > 256:
            raise ExperimentError("Idempotency key exceeds 256 characters")
        with self._transaction() as connection:
            row = self._sql(connection, "SELECT preview FROM platform_experiment_imports WHERE id=?", (preview_id,)).fetchone()
            if row is None:
                raise ExperimentError("Unknown result import preview")
            preview = json.loads(row[0])
            payload = self._batch(connection, preview["batch_id"], lock=True)
            prior = self._sql(connection, "SELECT preview_id,receipt FROM platform_experiment_receipts WHERE idempotency_key=?", (idempotency_key,)).fetchone()
            if prior:
                if prior[0] != preview_id:
                    raise ExperimentError("Idempotency key already belongs to a different import")
                return json.loads(prior[1])
            if not preview["valid"] or preview["batch_sha256"] != payload["payload_sha256"]:
                raise ExperimentError("Import preview is invalid or its frozen batch changed")
            raw = (self.artifact_root / preview["raw_sha256"]).read_bytes()
            if _sha(raw) != preview["raw_sha256"]:
                raise ExperimentError("Original upload failed integrity verification")
            existing = {r[0] for r in self._sql(connection, "SELECT observation_key FROM platform_experiment_observations WHERE batch_id=?", (preview["batch_id"],)).fetchall()}
            # Reparse the preserved original under the same policy. A modified
            # preview or concurrently confirmed upload cannot bypass validation.
            validated = validate_observations(parse_result_upload(raw, preview["filename"]), payload, existing_keys=existing)
            if not validated["valid"]:
                raise ExperimentError("Import conflicts with an existing observation or no longer validates")
            if canonical_json(validated["rows"]) != canonical_json(preview["rows"]):
                raise ExperimentError("Preview rows differ from the preserved original upload")
            for item in validated["rows"]:
                self._sql(connection, "INSERT INTO platform_experiment_observations VALUES (?,?,?,?,?)",
                          (preview["batch_id"], item["observation_key"], item["normalized"]["candidate_id"], preview_id, canonical_json(item)))
            receipt = {"preview_id": preview_id, "batch_id": preview["batch_id"],
                       "imported_count": len(validated["rows"]), "raw_sha256": preview["raw_sha256"],
                       "batch_sha256": payload["payload_sha256"], "confirmed_at": _now(), "status": "confirmed"}
            self._sql(connection, "INSERT INTO platform_experiment_receipts VALUES (?,?,?,?)",
                      (idempotency_key, preview_id, preview["batch_id"], canonical_json(receipt)))
            return receipt

    def observations(self, batch_id: str) -> list[dict]:
        with self._transaction() as connection:
            payload = self._batch(connection, batch_id)
            rows = self._sql(connection, "SELECT o.payload,o.preview_id,i.preview FROM platform_experiment_observations o JOIN platform_experiment_imports i ON i.id=o.preview_id WHERE o.batch_id=? ORDER BY o.candidate_id,o.observation_key", (batch_id,)).fetchall()
        expected_by_import: dict[str, dict | None] = {}
        results = []
        for row, preview_id, saved_preview in rows:
            item = json.loads(row)
            if preview_id not in expected_by_import:
                try:
                    preview = json.loads(saved_preview)
                    raw = (self.artifact_root / preview["raw_sha256"]).read_bytes()
                    if _sha(raw) != preview["raw_sha256"] or preview["batch_sha256"] != payload["payload_sha256"]:
                        raise ExperimentError("Original observation evidence failed integrity verification")
                    validated = validate_observations(parse_result_upload(raw, preview["filename"]), payload)
                    if not validated["valid"]:
                        raise ExperimentError("Original observation evidence no longer validates")
                    expected_by_import[preview_id] = {r["observation_key"]: r for r in validated["rows"]}
                except (OSError, ValueError, KeyError):
                    expected_by_import[preview_id] = None
            expected = (expected_by_import[preview_id] or {}).get(item.get("observation_key"))
            if expected is not None and canonical_json(expected) == canonical_json(item):
                results.append(item)
            else:
                results.append({"observation_key": item.get("observation_key"), "preview_id": preview_id,
                                "normalized": {"candidate_id": item.get("normalized", {}).get("candidate_id"),
                                               "target_id": payload["target_id"], "value": None,
                                               "kind": "observation_status", "status": "integrity_failed"},
                                "integrity_status": "quarantined", "raw": None,
                                "reason": "Stored observation or original upload failed integrity verification"})
        return results
