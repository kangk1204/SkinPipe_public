"""Whole-run offline bundles and split review workbooks, with leased publication."""

from __future__ import annotations

from datetime import datetime, timezone
from copy import deepcopy
import hashlib
import io
import json
import math
from pathlib import Path
import shutil
import tempfile
import zipfile

from .artifacts import canonical_json, sha256_file
from .engines import ResourceDeferred
from .identity import digest
from .export_assets import export_asset_root, implementation_files

REVIEW_FIELDS = [
    ("review_skin_score", "피부 적용 점수 (0–10점)"),
    ("review_permeation", "피부 투과 분류"),
    ("review_sensitisation", "감작성 구조 경보"),
    ("review_stability", "안정성 구조 경보"),
    ("review_phototoxicity", "광독성 구조 경보"),
    ("review_irritation", "피부 자극 구조 경보"),
    ("review_tox_Skin_Reaction", "피부 감작성 예측 (0–1)"),
    ("review_tox_AMES", "AMES 예측 (0–1)"),
    ("review_tox_hERG", "hERG 예측 (0–1)"),
    ("review_tox_DILI", "간 손상 예측 (0–1)"),
    ("review_tox_ClinTox", "임상 독성 예측 (0–1)"),
    ("review_tox_Carcinogens_Lagunin", "발암성 예측 (0–1)"),
    ("review_solubility", "수용해도 예측 (log mol/L)"),
    ("review_mw", "분자량 (g/mol)"),
    ("review_logp", "logP"),
    ("review_tpsa", "TPSA (Å²)"),
    ("review_hbd", "수소결합 공여체 수"),
    ("review_hba", "수소결합 수용체 수"),
    ("review_qed", "QED (0–1)"),
    ("review_admet_status", "ADMET 계산 상태"),
    ("review_admet_models", "ADMET 모델"),
    ("review_procurement", "구매 정보"),
]

METRIC_GUIDE = [
    ("priority_score_v2", "평가 점수", "0–100점", "↑ 높을수록 먼저 검토",
     "저장된 평가 기준에 따른 검토 우선순위입니다.", "같은 연구·계산 조건 안에서 비교",
     "SkinPipe 평가·선정 정책"),
    ("priority_rank_v2", "원분석 순위", "순위 / 전체 평가 수", "↓ 1위부터 검토",
     "13 / 10,000은 원분석 평가 대상 10,000개 중 13위입니다. 대조물질 포함 여부를 함께 봅니다.", "원분석의 평가 범위·정렬 기준 확인",
     "SkinPipe 저장 순위"),
    ("routine_review_order", "대표 후보 검토 순서", "순서 / 선정 수", "↓ 1번부터 검토",
     "선정 후보 안에서 확인할 순서입니다. 골격 다양성을 반영하므로 원분석 순위와 다를 수 있습니다.", "해당 연구·조건에서 선정한 후보 안에서 비교",
     "SkinPipe 대표 후보 선정 정책"),
    ("review_dock_affinity", "도킹", "kcal/mol", "↓ 더 음수일수록 유리",
     "예: −9는 −6보다 도킹 점수가 낮습니다. 먼저 실행 상태와 포즈 QC를 확인합니다.", "같은 수용체·포켓·엔진·준비·검색 조건에서 비교",
     "https://autodock-vina.readthedocs.io/en/latest/docking_basic.html"),
    ("review_skin_score", "피부 적용 점수", "0–10점", "↑ 높을수록 물성 기준 적합",
     "분자량·logP·수소결합 공여체·TPSA를 조합한 피부 적용 물성 점수입니다.", "같은 피부 평가 정책·조건에서 비교",
     "https://github.com/kangk1204/SkinPipe_public/blob/main/skinpipe/skinpipe/filters/topical.py"),
    ("review_permeation", "피부 투과 분류", "높음·중간·낮음", "목표 부위에 따라 판단",
     "높음은 더 잘 투과하는 방향입니다. 피부 표면 잔류와 깊은 부위 전달은 목표가 다릅니다.", "연구 목적·목표 부위·제형과 함께 검토",
     "SkinPipe 피부 투과 모델"),
    ("review_sensitisation", "감작성 구조 경보", "경보 목록", "경보 내용·수준 검토",
     "반응성 구조 등 감작성 검토가 필요한 특징을 확인합니다.", "구조 경보와 모델의 감작성 예측값을 나란히 검토",
     "SkinPipe 구조 경보 규칙"),
    ("review_stability", "안정성 구조 경보", "경보 목록", "경보 내용·수준 검토",
     "가수분해·산화 등 구조상의 안정성 검토 항목입니다.", "제형·보관·실험 조건과 함께 검토",
     "SkinPipe 구조 경보 규칙"),
    ("review_phototoxicity", "광독성 구조 경보", "없음·낮음·중간·높음", "경보 수준이 낮은 쪽 검토",
     "빛에 반응할 수 있는 구조 특징에 대한 경보입니다.", "노출 부위·광 노출 조건과 함께 검토",
     "SkinPipe 광반응 구조 규칙"),
    ("review_irritation", "피부 자극 구조 경보", "없음·낮음·중간·높음", "경보 수준이 낮은 쪽 검토",
     "산·염기성 및 자극 관련 구조 특징을 확인합니다.", "농도·pH·제형과 함께 검토",
     "SkinPipe 피부 자극 구조 규칙"),
]
for _endpoint, _label, _meaning in [
    ("Skin_Reaction", "피부 감작성 예측", "감작성 반응의 양성 방향을 예측하는 출력입니다."),
    ("AMES", "AMES 예측", "변이원성 양성 방향을 예측하는 출력입니다."),
    ("hERG", "hERG 예측", "hERG 채널 차단 방향을 예측하는 출력입니다."),
    ("DILI", "간 손상 예측", "약물 유발 간 손상 양성 방향을 예측하는 출력입니다."),
    ("ClinTox", "임상 독성 예측", "임상 독성 양성 방향을 예측하는 출력입니다."),
    ("Carcinogens_Lagunin", "발암성 예측", "발암성 양성 방향을 예측하는 출력입니다."),
]:
    METRIC_GUIDE.append(("review_tox_" + _endpoint, _label, "모델 출력 0–1", "↓ 낮을수록 독성 신호 작음",
                         _meaning, "ADMET-AI/TDC 양성 정의 기준. 같은 모델·지표끼리 비교. 다른 모델은 지표 정의 확인",
                         "https://tdcommons.ai/single_pred_tasks/tox/"))
METRIC_GUIDE.extend([
    ("review_solubility", "수용해도 예측", "AqSolDB: log(mol/L)", "↑ 높을수록 수용해도 증가",
     "예: −2는 −4보다 물에 더 잘 녹는 방향입니다.", "수용해도 증가 자체와 연구 적합성을 구분하고 제형 조건 확인",
     "https://github.com/swansonk14/admet_ai/blob/main/admet_ai/resources/data/admet.csv"),
    ("review_mw", "분자량", "g/mol", "크기·범위 검토",
     "분자의 크기 지표입니다. 작은 값만을 목표로 후보를 고르지 않습니다.", "피부 투과·구조·다른 물성과 함께 검토",
     "https://www.rdkit.org/docs/source/rdkit.Chem.Descriptors.html"),
    ("review_logp", "logP", "지용성 log 비율", "적정 범위 검토",
     "높을수록 지용성이 커집니다. 수용성과 피부 지질층 분배의 균형을 봅니다.", "연구 목적·제형과 함께 검토",
     "https://www.rdkit.org/docs/source/rdkit.Chem.Crippen.html"),
    ("review_tpsa", "TPSA", "Å²", "극성·범위 검토",
     "높을수록 극성 표면적이 큽니다.", "투과·용해성 등 다른 물성과 함께 검토",
     "https://www.rdkit.org/docs/source/rdkit.Chem.rdMolDescriptors.html"),
    ("review_hbd", "수소결합 공여체 수", "개수", "개수·범위 검토",
     "수소결합 공여에 참여할 수 있는 작용기 수입니다.", "개수 하나로 우열을 정하지 않고 다른 물성과 함께 검토",
     "https://www.rdkit.org/docs/source/rdkit.Chem.rdMolDescriptors.html"),
    ("review_hba", "수소결합 수용체 수", "개수", "개수·범위 검토",
     "수소결합 수용에 참여할 수 있는 작용기 수입니다.", "개수 하나로 우열을 정하지 않고 다른 물성과 함께 검토",
     "https://www.rdkit.org/docs/source/rdkit.Chem.rdMolDescriptors.html"),
    ("review_qed", "QED", "0–1", "↑ 높을수록 약물 유사성",
     "약물의 물성 분포를 바탕으로 계산한 약물 유사성 지표입니다.", "피부 적용 목적은 피부 평가 열에서 별도로 검토",
     "https://www.rdkit.org/docs/source/rdkit.Chem.QED.html"),
])
_GUIDE_BY_FIELD = {row[0]: row for row in METRIC_GUIDE}


def review_header(field, title):
    guidance = _GUIDE_BY_FIELD.get(field)
    return title + "\n" + guidance[3] if guidance else title


def add_metric_guide(book):
    from openpyxl.styles import Alignment, Font
    guide = book.create_sheet("지표 읽는 법", 1)
    guide.append(["점수·계산값 읽는 법"])
    guide.append(["비교 순서", "선정·제외·보류 상태를 먼저 보고, 같은 연구·계산 조건의 후보를 비교합니다."])
    guide.append(["결측·모델", "미평가·실패·값 오류는 0이 아닙니다. 복수 모델은 각 모델의 지표 정의를 확인합니다."])
    guide.append(["척도", "독성 0–1 출력의 판정 기준은 모델 정의에서 확인합니다. 물성 값은 연구 목적과 제형에 맞춰 봅니다."])
    for row in range(2, 5):
        guide.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
    guide.append([])
    guide.append(["지표", "단위·범위", "읽는 방향", "의미·예시", "비교 조건", "출처"])
    for row in METRIC_GUIDE:
        guide.append(list(row[1:]))
    for column, width in zip("ABCDEF", (27, 25, 37, 68, 65, 78)):
        guide.column_dimensions[column].width = width
    guide.freeze_panes = "A7"
    guide.auto_filter.ref = "A6:F" + str(guide.max_row)
    for cells in guide:
        for cell in cells:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        guide.row_dimensions[cells[0].row].height = 58 if cells[0].row >= 7 else 34
    for cell in guide[6]:
        cell.font = Font(bold=True)
    return guide

FIELDS = [
    ("candidate_id", "후보 ID"),
    ("name", "화합물명"),
    ("image", "2D 구조"),
    ("smiles", "SMILES"),
    ("priority_score_v2", "평가 점수 (0–100점)"),
    ("review_dock_affinity", "도킹 (kcal/mol)"),
    ("review_dock_status", "도킹 실행 상태"),
    *REVIEW_FIELDS,
    ("image_kind", "이미지 근거"),
    ("mw", "분자량 원자료 (g/mol)"),
    ("logp", "logP 원자료"),
    ("procurement", "구매 근거"),
    ("tox_Skin_Reaction", "피부 감작성 예측 (모델 출력)"),
    ("tox_AMES", "AMES 예측 (모델 출력)"),
    ("tox_Carcinogens_Lagunin", "발암성 예측 (모델 출력)"),
    ("tox_ClinTox", "임상 독성 예측 (모델 출력)"),
    ("tox_DILI", "간 손상 예측 (모델 출력)"),
    ("tox_hERG", "hERG 예측 (모델 출력)"),
    ("admet", "ADMET"),
    ("additional_toxicity", "추가 로컬 모델 · 기존 ADMET와 별도"),
    ("primary_docking", "주 도킹 결과"),
    ("docking", "도킹 전체 기록"),
    ("admet_search_risk", "독성·피부독성 검색 판단"),
    ("priority_score_status", "우선순위 계산 범위"),
    ("priority_missing_channels", "우선순위 결측 채널"),
    ("discovery_score", "탐색 점수"),
    ("discovery_origin", "발굴 경로"),
    ("discovery_filter_reasons", "필터 제외 사유"),
    ("selection_reason", "선정 이유"),
    ("computation_status", "계산 상태"),
    ("campaign_computation_status", "이번 캠페인의 계산 상태"),
    ("preparation_diagnostic", "구조 준비 보류 근거"),
    ("role", "연구 내 역할"),
    ('routine_disposition', '대표 후보 선정·보류 사유'),
    ('routine_pose_qc', '대표 도킹 자세 QC'),
    ("skin", "피부 적용 구조 평가"),
    ("structural_admet", "구조 기반 안전성 평가"),
    ("structural_admet_context", "피부 평가 조건"),
    ("hop", "호핑 비교 근거"),
    ("reference_similarity", "질의 구조와 2D 유사도 근거"),
    ("reference_alignment", "질의 구조와 형태·약리 특징 비교"),
    ("novel_scaffold", "새 환상 코어 여부"),
    ("acyclic_topology_comparison", "비고리 구조 비교 근거"),
    ("target_association", "표적 활성 근거 존재 여부"),
]

DETAILS_PER_WORKBOOK = 200
STRUCTURE_ASSET_KEYS = {
    "observed_sdf_sha256", "original_structure_sha256", "protein_structure_sha256",
    "original_protein_heavy_coordinates_sha256", "retained_components_sha256",
    "assembly_structure_sha256", "assembly_metadata_sha256", "ccd_sha256",
    "template_sdf_sha256", "instance_sdf_sha256", "raw_sha256", "artifact_sha256", "ligand_sdf",
}


def review_columns(row):
    """Expose existing values without calculating predictions or merging models."""
    def mapping(value):
        return value if isinstance(value, dict) else {}

    def number(value):
        try:
            return value if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value) else None
        except OverflowError:
            return None

    predictions = [p for p in row.get("admet") or [] if isinstance(p, dict)]

    def values(endpoint, *, classifier=False):
        found, invalid = [], []
        for prediction in predictions:
            if prediction.get("status") not in {"success", "partial_nonfinite"}:
                continue
            payload = prediction.get("values")
            if payload is None and isinstance(prediction.get("values_json"), str):
                try:
                    payload = json.loads(prediction["values_json"])
                except ValueError:
                    payload = None
            raw = mapping(payload).get(endpoint)
            value = number(raw)
            if value is not None and (not classifier or 0 <= value <= 1):
                found.append((prediction.get("model_id") or "모델 미기록", value))
            elif raw is not None:
                invalid.append(str(prediction.get("model_id") or "모델 미기록"))
        if not found:
            return "값 오류 · 원자료 확인" if invalid else "미평가"
        if len(found) == 1 and not invalid:
            return found[0][1]
        return "\n".join([str(model) + ": " + repr(value) for model, value in found]
                         + [model + ": 값 오류 · 원자료 확인" for model in invalid])

    def alerts(skin, key, status):
        value = skin.get(key)
        if not isinstance(value, list) or skin.get(status) != "EVALUATED":
            return "미평가"
        return "; ".join(str(item) for item in value) if value else "경보 없음"

    skin = mapping(row.get("skin"))
    structural = mapping(row.get("structural_admet"))
    concern = {"none": "경보 없음", "low": "낮음", "moderate": "중간", "high": "높음", "critical": "매우 높음"}
    permeation = skin.get("permeation_class")
    permeation = {"good": "높음 (good)", "moderate": "중간 (moderate)", "poor": "낮음 (poor)"}.get(permeation, permeation or "미평가")
    primary = mapping(row.get("primary_docking"))
    observations = row.get("docking") or []
    observations = [observations] if isinstance(observations, dict) else observations
    index = primary.get("observation_index", 0)
    if observations and not any(key in primary for key in ("score", "value", "affinity")):
        primary = mapping(observations[index]) if isinstance(index, int) and 0 <= index < len(observations) else primary
    from .procurement import has_purchase_record
    fields = {
        "review_skin_score": number(skin.get("topical_score", row.get("topical_score"))),
        "review_permeation": permeation,
        "review_sensitisation": alerts(skin, "sensitiser_alerts", "sensitisation_status"),
        "review_stability": alerts(skin, "stability_alerts", "stability_status"),
        "review_phototoxicity": concern.get(mapping(structural.get("phototoxicity")).get("concern"), "미평가"),
        "review_irritation": concern.get(mapping(structural.get("skin_irritation")).get("concern"), "미평가"),
        "review_solubility": values("Solubility_AqSolDB"),
        "review_mw": number(row.get("mw")) if number(row.get("mw")) is not None else values("molecular_weight"),
        "review_logp": number(row.get("logp")) if number(row.get("logp")) is not None else values("logP"),
        "review_tpsa": values("tpsa"), "review_hbd": values("hydrogen_bond_donors"),
        "review_hba": values("hydrogen_bond_acceptors"), "review_qed": values("QED"),
        "review_admet_status": "\n".join(str(p.get("model_id") or "모델 미기록") + ": " +
            {"success": "계산 완료", "partial_nonfinite": "일부 결측", "failed": "실패", "unsupported_structure": "미지원"}.get(p.get("status"), str(p.get("status") or "미평가"))
            for p in predictions) or "미평가",
        "review_admet_models": "\n".join(str(p.get("model_id") or "모델 미기록") for p in predictions) or "미기록",
        "review_procurement": row.get("procurement_label") or ("구매 정보 있음" if has_purchase_record(row) else "구매 정보 미기록"),
        "review_dock_affinity": number(primary.get("score", primary.get("affinity", primary.get("value")))),
        "review_dock_status": primary.get("status") or "미평가",
    }
    for endpoint in ("Skin_Reaction", "AMES", "hERG", "DILI", "ClinTox", "Carcinogens_Lagunin"):
        fields["review_tox_" + endpoint] = values(endpoint, classifier=True)
    return {key: "미평가" if value is None else value for key, value in fields.items()}


def xlsx_value(value, *, candidate_id, field, warnings):
    value = text(value)
    invalid_number = False
    if type(value) in (int, float):
        try:
            invalid_number = not math.isfinite(value)
        except OverflowError:
            invalid_number = True
    if invalid_number or isinstance(value, str) and (
        len(value.encode("utf-16-le", errors="surrogatepass")) // 2 > 32767
        or any(not (ord(c) in (9, 10, 13) or 32 <= ord(c) <= 0xD7FF
                       or 0xE000 <= ord(c) <= 0xFFFD or 0x10000 <= ord(c) <= 0x10FFFF)
               for c in value)
    ):
        warnings.append({"candidate_id": candidate_id, "field": field,
                         "reason": "xlsx_value_requires_raw_chunk_reference"})
        # The identifier itself may contain the same unsupported characters.
        # Escape and bound this display reference; warnings and JSON retain it.
        reference = json.dumps({"candidate_id": candidate_id, "field": field},
                               ensure_ascii=True, default=str)
        return "원자료 파일 참조 · " + reference[:240]
    return value


def text(value):
    if value is None:
        return None
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, allow_nan=False)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        value = "'" + value
    return value


def toxicity_columns(row, warnings):
    """Keep raw predictions; quarantine malformed display values individually."""
    from .admet_search import RISK_ENDPOINTS
    columns = {endpoint: [] for endpoint in RISK_ENDPOINTS}
    for index, prediction in enumerate(row.get("admet") or []):
        if not isinstance(prediction, dict) or prediction.get("status") not in {"success", "partial_nonfinite"}:
            continue
        values = prediction.get("values")
        if values is None and isinstance(prediction.get("values_json"), str):
            try:
                values = json.loads(prediction["values_json"])
            except (ValueError, TypeError):
                values = None
        invalid = not isinstance(values, dict)
        if not invalid:
            invalid = any(
                value is not None and (isinstance(value, bool) or not isinstance(value, (int, float))
                    or not 0 <= value <= 1 or not math.isfinite(value))
                for endpoint in RISK_ENDPOINTS
                for value in [values.get(endpoint)]
            )
        if invalid:
            warnings.append({"candidate_id": row.get("candidate_id"), "prediction_index": index,
                             "component": "admet", "status": "quarantined",
                             "reason": "invalid_toxicity_display_values"})
            continue
        for endpoint in RISK_ENDPOINTS:
            if values.get(endpoint) is not None:
                columns[endpoint].append({"model_id": prediction.get("model_id"), "value": values[endpoint]})
    return {"tox_" + endpoint: values or None for endpoint, values in columns.items()}


class RunExports:
    def __init__(self, service):
        self.service = service

    @staticmethod
    def implementation_identity():
        files = {name: sha256_file(path) for name, path in implementation_files().items()}
        return {
            "version": "whole-run-export-v5",
            "files": files,
            "sha256": digest(files),
        }

    def submit(self, body, *, owner="researcher"):
        if not isinstance(body, dict) or set(body) - {
            "run_id",
            "run_ids",
            "snapshot_id",
            "target_id",
            "format",
            "retry_of",
            "candidate_details",
        }:
            raise ValueError("Export accepts a frozen run or snapshot and format")
        request = dict(body)
        if not isinstance(request.get("format", "zip"), str) or request.get("format", "zip") not in {"zip", "xlsx"}:
            raise ValueError("format must be zip or xlsx")
        if "candidate_details" in request and not isinstance(request["candidate_details"], bool):
            raise ValueError("candidate_details must be a boolean")
        request["candidate_details"] = request.get(
            "candidate_details", request.get("format", "zip") == "xlsx"
        )
        if "run_ids" in request:
            ids = request["run_ids"]
            if (not isinstance(ids, list) or not 1 <= len(ids) <= 200
                    or any(not isinstance(value, str) or not value or len(value) > 160 for value in ids)
                    or len(set(ids)) != len(ids)):
                raise ValueError("run_ids requires 1..200 distinct bounded run identifiers")
            if any(request.get(key) is not None for key in ("run_id", "snapshot_id", "target_id")):
                raise ValueError("Choose multiple runs or one frozen run/snapshot")
            request["runs"] = [self._freeze_run(value, owner) for value in ids]
        else:
            if request.get("run_id"):
                frozen = self._freeze_run(request["run_id"], owner)
                request.update(frozen)
            elif not request.get("snapshot_id"):
                raise ValueError("Export requires a frozen run or snapshot")
            if request.get("snapshot_id"):
                manifest = self.service.catalog.manifest(request["snapshot_id"])
                if self.service.catalog.snapshot_owner(request["snapshot_id"]) != owner:
                    raise PermissionError("Export snapshot belongs to another project")
                if request.get("target_id") not in manifest["target_counts"] and manifest["candidate_count"]:
                    raise ValueError("Export requires a target from the frozen snapshot")
                request["manifest_sha256"] = digest(manifest)
        request.update(
            format=request.get("format", "zip"),
            exporter=self.implementation_identity(),
        )
        retry_of = request.get("retry_of")
        if retry_of is not None:
            previous = self.service.queue.get(retry_of, owner=owner)
            original = previous.get("request", {})
            def population(value):
                return [{k: item.get(k) for k in ("run_id", "snapshot_id", "target_id", "manifest_sha256", "run_request")}
                        for item in value.get("runs", [])]
            if (previous["kind"] != "result_export"
                    or previous["status"] not in {"failed", "budget_deferred", "cancelled"}
                    or any(original.get(k) != request.get(k) for k in (
                        "snapshot_id", "target_id", "manifest_sha256", "format",
                        "run_ids")) or population(original) != population(request)
                    or original.get("candidate_details", original.get("format", "zip") == "xlsx") != request["candidate_details"]):
                raise ValueError("Export retry requires a terminal export of the same frozen population")
            if "runs" in original:
                request["runs"] = deepcopy(original["runs"])
            for key in ("run_summary", "run_request"):
                if key in original:
                    request[key] = deepcopy(original[key])
            request["retry_history"] = {
                "previous_export_id": retry_of,
                "previous_status": previous["status"],
                "previous_compute_seconds": previous.get("compute_seconds"),
                "policy": "explicit new bounded export; previous job, consumption and failures preserved",
            }
        signature = digest(request)
        job = self.service.queue.enqueue(
            "result_export",
            request,
            owner=owner,
            project="exports-" + digest(owner)[:24],
            idempotency_key="export-" + signature,
            budget_seconds=3600,
        )
        return {
            "export_id": job["id"],
            "status": job["status"],
            "scope": "all saved target candidates including filter exclusions",
            "container_format": "zip",
            "requested_format": request["format"],
            "retry_of": retry_of,
        }

    def _freeze_run(self, run_id, owner):
        run = deepcopy(self.service.discovery.get(run_id, owner=owner))
        record = self.service.discovery._record(run_id, owner) if hasattr(self.service.discovery, "_record") else {}
        snapshot_id = run.get("result_snapshot_id") or run.get("snapshot_id")
        result = {"run_id": run_id, "run_summary": run, "run_request": deepcopy(record.get("request", run.get("request", {}))),
                  "snapshot_id": snapshot_id, "target_id": run.get("target_id"), "candidate_count": 0}
        if snapshot_id:
            source = self.service.catalog.manifest(snapshot_id)
            if self.service.catalog.snapshot_owner(snapshot_id) != owner:
                raise PermissionError("Export snapshot belongs to another project")
            target = result["target_id"]
            if target not in source["target_counts"] and source["candidate_count"]:
                raise ValueError("Run target differs from the published snapshot")
            result.update(manifest_sha256=digest(source), candidate_count=source["target_counts"].get(target, 0))
        return result

    def status(self, export_id, *, owner="researcher"):
        job = self.service.queue.get(export_id, owner=owner)
        if job["kind"] != "result_export":
            raise KeyError("Unknown result export")
        return {
            "export_id": export_id,
            "status": job["status"],
            "progress": job.get("progress"),
            "result": job.get("result"),
            "retryable_as_new_export": job["status"] in {"failed", "budget_deferred", "cancelled"},
            "retry_policy": "explicit retry_of submission creates a separate 3600-second allowance; no automatic reset",
        }

    def download(self, export_id, *, owner="researcher"):
        result = self.status(export_id, owner=owner)
        if result["status"] != "complete":
            raise ValueError("Export has not completed")
        receipt = result["result"]
        path = self.service.artifacts.verified_path(receipt["sha256"])
        return path, "application/zip", receipt["filename"]

    def execute(self, request, *, stop, progress):
        if (
            request.get("exporter") is not None
            and request["exporter"] != self.implementation_identity()
        ):
            raise ValueError(
                "Frozen export implementation changed; original request preserved"
            )
        catalog, artifacts = self.service.catalog, self.service.artifacts
        sources = request.get("runs", [request])
        population, source_assets = 0, 0
        for entry in sources:
            if not entry.get("snapshot_id"):
                continue
            source = catalog.manifest(entry["snapshot_id"])
            if digest(source) != entry["manifest_sha256"]:
                raise ValueError("Export source manifest changed")
            population += source.get("target_counts", {}).get(entry.get("target_id"), source.get("candidate_count", 0))
            source_assets += self._structure_asset_bytes(source, stop=stop)
        # Separate run directories deliberately retain separate observations.
        # Large original coordinates therefore occupy space once per run,
        # followed by the ZIP and its CAS publication copy.
        required = max(2 * 1024**3, population * 400000 + 3 * source_assets)
        if shutil.disk_usage(self.service.root).free < required:
            raise ResourceDeferred("export_resource_deferred")
        root = self.service.root / "export-work"
        root.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=root, prefix="export-") as directory:
            directory = Path(directory)
            receipt = self.build(request, directory, stop=stop, progress=progress)
            output = directory / "results.zip"
            stop()
            with zipfile.ZipFile(
                output, "w", zipfile.ZIP_DEFLATED, compresslevel=1
            ) as bundle:
                for path in sorted(directory.rglob("*")):
                    stop()
                    if path.is_file() and path != output:
                        bundle.write(path, str(path.relative_to(directory)))
            artifact = artifacts.put_file(output, media_type="application/zip")
            return {
                **artifact,
                "filename": "SkinPipe-"
                + (request.get("snapshot_id") or ("runs-" + digest(request.get("run_ids", [request.get("run_id")]))[:24]))
                + "-"
                + request["format"]
                + ".zip",
                "manifest": receipt,
                "candidate_count": receipt["candidate_count"],
            }

    def _structure_asset_bytes(self, source, *, stop):
        """Bound the added coordinate-copy cost without reading global DBs."""
        descriptor_sha = source.get("provenance", {}).get("target_descriptor_sha256")
        if not isinstance(descriptor_sha, str):
            return 0
        stop()
        artifacts = self.service.artifacts
        try:
            raw = artifacts.get_bytes(descriptor_sha)
            total = len(raw)
            if len(raw) > 16 * 1024**2:
                return total
            descriptor = json.loads(raw)
        except (ValueError, TypeError, OSError):
            return 0  # the normal exporter preserves its missing-asset warning
        pending, seen = [descriptor], {descriptor_sha}
        while pending:
            stop()
            value = pending.pop()
            if isinstance(value, list):
                pending.extend(value)
            elif isinstance(value, dict):
                for key, item in value.items():
                    if key in STRUCTURE_ASSET_KEYS and isinstance(item, str) and item not in seen:
                        seen.add(item)
                        try:
                            total += artifacts.verified_path(item).stat().st_size
                        except (ValueError, OSError):
                            pass
                    elif isinstance(item, (dict, list)):
                        pending.append(item)
        return total

    def build(self, request, directory, *, stop=lambda: None, progress=lambda _: None):
        directory = Path(directory)
        directory.mkdir(parents=True, exist_ok=True)
        if "runs" in request:
            return self._build_many(request, directory, stop=stop, progress=progress)
        if not request.get("snapshot_id"):
            return self._build_empty(request, directory)
        with (directory / "candidates.sdf").open("w", encoding="utf-8") as sdf:
            return self._build(
                request, directory, sdf=sdf, stop=stop, progress=progress
            )

    def _build_empty(self, request, directory):
        from openpyxl import Workbook
        book = Workbook()
        summary = book.active
        summary.title = "실행 요약"
        warnings = []
        run = request.get("run_summary", {})
        for label, value in [("실행 ID", request.get("run_id")), ("저장 후보 수", 0),
                             ("실행 상태", run.get("status")), ("보류·실패 사유", run.get("zero_reasons", run.get("reason"))),
                             ("저장 범위", "내보내기 요청 시점에 공개된 후보 스냅샷 없음"),
                             ("과학적 검증", "NOT_ESTABLISHED")]:
            summary.append([label, xlsx_value(value, candidate_id="run", field=label, warnings=warnings)])
        conditions = book.create_sheet("입력·조건")
        for key, value in request.get("run_request", {}).items():
            conditions.append([text(key), xlsx_value(value, candidate_id="run", field=key, warnings=warnings)])
        for sheet in book:
            sheet.column_dimensions["A"].width = 28
            sheet.column_dimensions["B"].width = 110
        book.save(directory / "review-0001.xlsx")
        (directory / "run-summary.json").write_bytes(canonical_json(run))
        (directory / "run-conditions.json").write_bytes(canonical_json(request.get("run_request", {})))
        (directory / "chunks").mkdir()
        (directory / "chunks/candidates-00001.json").write_bytes(canonical_json([]))
        receipt = {"schema": "skinpipe-run-export-v1", "run_id": request.get("run_id"), "snapshot_id": None,
            "exporter": request.get("exporter", self.implementation_identity()),
            "container_format": "zip", "requested_format": request.get("format", "zip"),
            "retry_history": request.get("retry_history"),
            "candidate_count": 0, "candidate_ids_in_order": [], "candidate_order_sha256": digest([]),
            "workbooks": ["review-0001.xlsx"], "candidate_detail_tabs": [],
            "chunks": [{"file": "chunks/candidates-00001.json", "count": 0, "sha256": digest([])}],
            "scope": "frozen run status and original conditions; no candidate snapshot at export submission",
            "publication_scope": {"status": run.get("status"), "whole_workflow_complete": False},
            "warnings": warnings, "scientific_validation": "NOT_ESTABLISHED",
            "files": {str(p.relative_to(directory)): {"sha256": sha256_file(p), "size_bytes": p.stat().st_size}
                      for p in directory.rglob("*") if p.is_file()}}
        (directory / "manifest.json").write_bytes(canonical_json(receipt))
        return receipt

    def _build_many(self, request, directory, *, stop, progress):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        book = Workbook()
        sheet = book.active
        sheet.title = "실행 요약"
        sheet.append(["실행 ID", "입력", "실행 상태", "계산 대상", "저장 행 수 · 질의·대조 포함", "후보 Excel", "원자료 폴더"])
        sheet.freeze_panes = "A2"
        runs, warnings = [], []
        for index, entry in enumerate(request["runs"], 1):
            stop()
            relative = "runs/" + str(index).zfill(4) + "-" + digest(entry["run_id"])[:16]
            child_request = {**entry, **{k: request[k] for k in ("format", "candidate_details", "exporter") if k in request}}
            child = self.build(child_request, directory / relative, stop=stop,
                progress=lambda value: progress({**value, "run_index": index, "run_count": len(request["runs"])}))
            run = entry.get("run_summary", {})
            record = {"run_id": entry["run_id"], "directory": relative, "snapshot_id": entry.get("snapshot_id"),
                      "target_id": entry.get("target_id"), "candidate_count": child["candidate_count"],
                      "run_status": run.get("status"), "workbooks": [relative + "/" + name for name in child["workbooks"]],
                      "manifest": relative + "/manifest.json"}
            runs.append(record)
            values = [entry["run_id"], run.get("research_input", entry.get("run_request", {}).get("input")),
                      run.get("status"), entry.get("target_id"), child["candidate_count"], record["workbooks"][0], relative]
            sheet.append([xlsx_value(value, candidate_id=entry["run_id"], field=str(i), warnings=warnings) for i, value in enumerate(values)])
            sheet.cell(sheet.max_row, 6).hyperlink = record["workbooks"][0]
            sheet.cell(sheet.max_row, 6).style = "Hyperlink"
        sheet.auto_filter.ref = sheet.dimensions
        for cells in sheet:
            for cell in cells:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column, width in zip("ABCDEFG", (44, 70, 24, 48, 18, 70, 60)):
            sheet.column_dimensions[column].width = width
        book.save(directory / "SkinPipe-all-runs.xlsx")
        (directory / "frozen-runs.json").write_bytes(canonical_json(request["runs"]))
        (directory / "읽어주세요.txt").write_text("SkinPipe-all-runs.xlsx에서 실행별 Excel을 여세요. 각 runs 폴더에 전체 후보, 제외 후보, 원자료, 입력 조건을 보존했습니다. 서로 다른 실행의 점수나 관측값은 합치지 않았습니다.\n", encoding="utf-8")
        receipt = {"schema": "skinpipe-multi-run-export-v1", "run_count": len(runs),
            "candidate_count": sum(r["candidate_count"] for r in runs), "runs": runs,
            "workbooks": ["SkinPipe-all-runs.xlsx"], "scope": "separate frozen runs; every saved candidate and exclusion retained per run",
            "observations_merged": False, "warnings": warnings, "scientific_validation": "NOT_ESTABLISHED",
            "files": {str(p.relative_to(directory)): {"sha256": sha256_file(p), "size_bytes": p.stat().st_size}
                      for p in directory.rglob("*") if p.is_file()}}
        (directory / "manifest.json").write_bytes(canonical_json(receipt))
        return receipt

    def _build(self, request, directory, *, sdf, stop, progress):
        from rdkit import Chem
        from rdkit.Chem import Draw, rdDepictor
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image
        from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils.units import pixels_to_EMU
        from .pose_view import candidate_docking_view

        source = self.service.catalog.manifest(request["snapshot_id"])
        target = request["target_id"]
        directory = Path(directory)
        (directory / "chunks").mkdir()
        (directory / "assets").mkdir()
        (directory / "views").mkdir()
        (directory / "images").mkdir()
        ids, chunks, books, warnings, count, chunk = [], [], [], [], 0, []
        workbook, sheet = None, None
        endpoint_metadata, additional_endpoint_metadata = {}, {}
        detail_tabs = []
        with_details = request.get("candidate_details", request.get("format") == "xlsx")
        workbook_limit = DETAILS_PER_WORKBOOK if with_details else 10000
        provenance = source.get("provenance", {})
        run_summary = request.get("run_summary")
        if run_summary is not None:
            (directory / "run-summary.json").write_bytes(canonical_json(run_summary))
            (directory / "run-conditions.json").write_bytes(canonical_json(request.get("run_request", {})))

        def new_book():
            book = Workbook()
            table = book.active
            table.title = "후보 검토"
            table.append([review_header(field, title) for field, title in FIELDS])
            table.freeze_panes = "A2"
            table.row_dimensions[1].height = 64
            for cell in table[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            from openpyxl.utils import get_column_letter

            for i in range(len(FIELDS)):
                table.column_dimensions[get_column_letter(i + 1)].width = 24
            for i, (field, _) in enumerate(FIELDS, 1):
                if field in {"priority_score_v2", "review_skin_score", "review_logp", "review_tpsa", "review_mw"} or field.startswith("review_tox_"):
                    table.column_dimensions[get_column_letter(i)].width = 20
            table.column_dimensions["C"].width = 36
            table.column_dimensions["D"].width = 60
            return book, table

        def add_structure_image(table, path, *, row, column, width=240, height=140):
            image = Image(str(path))
            image.width, image.height = width, height
            # Tie both corners to the candidate's cell rather than an absolute
            # drawing position.
            image.anchor = TwoCellAnchor(
                editAs="twoCell",
                _from=AnchorMarker(col=column, row=row - 1,
                                   colOff=pixels_to_EMU(6), rowOff=pixels_to_EMU(6)),
                to=AnchorMarker(col=column, row=row - 1,
                                colOff=pixels_to_EMU(width + 6), rowOff=pixels_to_EMU(height + 6)),
            )
            table.add_image(image)

        def flush_book():
            if workbook is None:
                return
            sheet.auto_filter.ref = sheet.dimensions
            add_metric_guide(workbook)
            info = workbook.create_sheet("원자료 계약")
            for row in [
                ("스냅샷", request["snapshot_id"]),
                ("실행 ID", request.get("run_id")),
                ("실행 상태", xlsx_value((run_summary or {}).get("status"), candidate_id="run", field="status", warnings=warnings)),
                ("범위", "전체 저장 후보 · 필터 제외 후보 포함"),
                ("결측", "빈 수치는 미평가·미기록이며 0이 아닙니다."),
                ("구조", "후보 표와 개별 탭의 2D 구조는 입력 SMILES에서 생성. 저장 도킹 포즈는 개별 탭에 별도로 표시"),
                ("계산값 열", "피부 평가 → 주요 독성 예측 → 물성 → 모델·구매 정보. 원래 ADMET 전체 값은 뒤쪽 원자료 열에 보존"),
                ("ADMET 표시", "단일 모델 값은 숫자로 표시. 복수 모델은 모델별 값을 구분. 미평가·실패는 0으로 바꾸지 않음"),
                ("점수", "우선순위 휴리스틱 · 효능이나 결합 확률 아님"),
                ("도킹", "같은 수용체·포켓·엔진·주 프로토콜 비교군 내 해석"),
                ("개별 화합물 탭", "C번호 탭 · 전체 후보 및 제외 후보 포함" if with_details else "미요청"),
                ("입력 식별", xlsx_value(provenance.get("research_input", provenance.get("input", provenance.get("target_descriptor"))), candidate_id="input", field="research_input", warnings=warnings)),
                ("피부 평가 조건", xlsx_value(provenance.get("skin_context"), candidate_id="input", field="skin_context", warnings=warnings)),
                ("일반 DB 다양성 탐색", xlsx_value(provenance.get("catalog_selection"), candidate_id="input", field="catalog_selection", warnings=warnings)),
                ("분석 완료 범위", json.dumps(source.get("provenance", {}).get("publication_scope", {
                    "status": "NOT_RECORDED", "whole_workflow_complete": None
                }), ensure_ascii=False)),
                ("우선순위 범위", "재계산 풀 밖의 보존 점수는 계산 범위 열로 구분. 서로 다른 풀의 점수를 직접 비교하지 않습니다."),
                (
                    "긴 원자료",
                    "전체 값은 chunks/*.json에 보존. Excel 셀 32767자 초과는 해당 파일 참조",
                ),
            ]:
                info.append(row)
            info.column_dimensions["A"].width = 16
            info.column_dimensions["B"].width = 100
            for row in info:
                for cell in row:
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                info.row_dimensions[row[0].row].height = 30
            if endpoint_metadata:
                definitions = workbook.create_sheet("ADMET 지표")
                definitions.append(["모델", "지표", "이름", "단위", "종", "측정 여부"])
                for model, items in sorted(endpoint_metadata.items()):
                    for item in items:
                        definitions.append(
                            [
                                xlsx_value(value, candidate_id=model,
                                           field="endpoint_" + field, warnings=warnings)
                                for field, value in (
                                    ("model", model),
                                    ("id", item.get("id") or item.get("endpoint")),
                                    ("name", item.get("name")),
                                    ("units", item.get("units", item.get("unit"))),
                                    ("species", item.get("species")),
                                )
                            ] + [
                                "모델 예측 · 실측 아님",
                            ]
                        )
                for column, width in zip("ABCDEF", (48, 32, 36, 22, 14, 30)):
                    definitions.column_dimensions[column].width = width
                definitions.freeze_panes = "A2"
                definitions.auto_filter.ref = definitions.dimensions
                for row in definitions:
                    for cell in row:
                        cell.alignment = Alignment(vertical="center", wrap_text=True)
                    definitions.row_dimensions[row[0].row].height = 30
                for cell in definitions[1]:
                    cell.font = Font(bold=True)
            name = "review-" + str(len(books) + 1).zfill(4) + ".xlsx"
            workbook.save(directory / name)
            books.append(name)

        def candidate_sheet(row, image_path, pose_image_path):
            title = "C" + str(count).zfill(6)
            detail = workbook.create_sheet(title)
            detail.append(["화합물 연구 기록", text(row["candidate_id"])])
            detail.append(["전체 후보 표", "후보 검토로 돌아가기"])
            detail["B2"].hyperlink = "#'후보 검토'!A" + str(sheet.max_row)
            detail["B2"].style = "Hyperlink"
            detail.append(["근거 범위", "계산·구조 평가이며 피부 개선 효능의 실측 근거가 아닙니다."])
            detail["C3"] = "읽는 방향"
            def brief(value, missing="미평가·미기록"):
                if value is None or value == [] or value == {} or value == "":
                    return missing
                if isinstance(value, dict):
                    return " · ".join(str(value[key]) for key in ("name", "severity", "status", "reason", "verdict") if value.get(key) is not None) or "상세 원자료 참조"
                if isinstance(value, (list, tuple)):
                    return "; ".join(brief(item) for item in value[:8]) + (" · 추가 항목은 원자료 참조" if len(value) > 8 else "")
                if isinstance(value, float):
                    return f"{value:.6g}"
                return str(value)

            def add(label, value):
                value = brief(value)
                if len(value) > 800:
                    value = value[:760] + "… 전체 값은 원자료 참조"
                detail.append([label, xlsx_value(value, candidate_id=row["candidate_id"], field=label, warnings=warnings)])
                field = next((key for key, title in FIELDS if title == label), None)
                if field in _GUIDE_BY_FIELD:
                    detail.cell(detail.max_row, 3, _GUIDE_BY_FIELD[field][3])

            def mapping(value):
                return value if isinstance(value, dict) else {}
            context = {**mapping(provenance.get("skin_context")), **mapping(row.get("structural_admet_context"))}
            skin = mapping(row.get("skin"))
            structural = mapping(row.get("structural_admet"))
            hop = mapping(row.get("hop"))
            primary = mapping(row.get("primary_docking"))
            observations = row.get("docking") or []
            observations = [observations] if isinstance(observations, dict) else observations
            if observations and "score" not in primary:
                index = primary.get("observation_index", 0)
                primary = observations[index] if isinstance(index, int) and 0 <= index < len(observations) else primary
            qualification = mapping(provenance.get("bound_complex_docking"))
            redocking = mapping(qualification.get("redocking"))
            research_input = mapping(provenance.get("research_input") or (request.get("run_request") or {}).get("input"))
            add("화합물명", row.get("name"))
            add("평가 점수 (0–100점)", row.get("priority_score_v2"))
            role = {"query_seed": "질의 구조 · 실측 활성 대조가 아님", "candidate": "탐색 후보", "control": "대조 구조"}.get(row.get("role", "candidate"), row.get("role"))
            add("연구 내 역할", role)
            routine = mapping(row.get('routine_disposition'))
            states = {'selected': '대표 후보', 'alternative': '검토 대안', 'held': '보류',
                      'pending': '평가 대기', 'duplicate': '중복 구조', 'control': '출발구조·대조물질'}
            reason_labels = {'saved_priority': '연구 내 저장 우선순위',
                'scaffold_or_acyclic_topology_diversity': '골격·비고리 연결 구조 다양성',
                'eligible_priority_backfill': '적격 후보 우선순위 보충',
                'representative_limit': '대표 후보 수 한도',
                'explicit_stereoisomer_required': '입체이성질체 지정 필요',
                'exact_pose_QC_required': '해당 도킹 자세 QC 필요',
                'pose_QC_failed': '도킹 자세 QC 기준 미달',
                'docking_score_review_required': '도킹 점수·실행 검토 필요',
                'docking_execution_incomplete': '도킹 실패·미완료',
                'exact_model_predictions_required': '정확 구조의 모델 예측 필요',
                'skin_structure_assessment_required': '피부·구조 평가 필요',
                'structural_alert_review_required': '구조 경보 검토',
                'receptor_or_calculation_model_held': '수용체·계산 모델 보류',
                'severe_predicted_admet_review_required': 'ADMET 경보 검토',
                'saved_priority_required': '연구 내 우선순위 필요'}
            if routine:
                add('대표 후보 선정·보류 사유', states.get(routine.get('state'), routine.get('state'))
                    + ' · ' + ' / '.join(reason_labels.get(r, r) for r in routine.get('reasons', [])))
            add('대표 도킹 자세 QC', row.get('routine_pose_qc'))
            add("연구 실행", request.get("run_id") or row.get("discovery_run_id") or provenance.get("discovery_run_id"))
            add("원래 연구 입력", (str(research_input.get("kind", "")) + " / " + str(research_input.get("format", "")) + " : " + str(research_input.get("value", ""))) if research_input else None)
            add("실제 계산 대상", row.get("target_id") or target)
            catalog_selection = mapping(provenance.get("catalog_selection"))
            if catalog_selection.get("fallback_used") is True:
                add("일반 DB 후보 선정", "구조 다양성 부분집합 · 선택 단계의 현재 단백질 활성·결합 미평가")
                add("DB 후보 선정 범위", "전체 인덱스 " + brief(catalog_selection.get("source_population"))
                    + " · 비교 풀 " + brief(catalog_selection.get("source_pool_size"))
                    + " · 확보 " + brief(catalog_selection.get("selected_count"))
                    + " · 방식 " + brief(catalog_selection.get("strategy")))
            add("계산·선정 상태", row.get("computation_status") or row.get("campaign_computation_status") or ("필터 통과" if row.get("discovery_filter_pass") is True else "필터 제외" if row.get("discovery_filter_pass") is False else None))
            add("선정·제외 근거", row.get("selection_reason") or row.get("discovery_filter_reasons"))
            add("구조 준비·도킹 보류", row.get("preparation_diagnostic") or qualification.get("hold_reasons") or (row.get("docking_execution") or {}).get("reason"))
            add("SMILES · 원래 입력 구조", row.get("smiles"))
            for field, label in REVIEW_FIELDS:
                add(label, row.get(field))
            add("분자 특성", "분자량 " + brief(row.get("mw")) + " g/mol · logP " + brief(row.get("logp")))
            add("연구 목적", context.get("research_goal"))
            add("피부 평가 조건", "제품 " + brief(context.get("product")) + " · 함량 " + brief(context.get("applied_percent")) + ("%" if context.get("applied_percent") is not None else "")
                + " · 부위 " + brief(context.get("application_site")) + " · 전달 목표 " + brief(context.get("goal")))
            add("피부 구조 평가", "국소 적용 점수 " + brief(skin.get("topical_score", row.get("topical_score")))
                + " · 투과 대리 지표 " + brief(skin.get("permeation_class")) + " · 실제 전달 깊이·효능 미검증")
            def alert_summary(alerts, status):
                return "규칙 평가에서 경고 기록 없음 · 안전성 미확립" if alerts == [] and status == "EVALUATED" else brief(alerts)
            add("감작성·안정성 구조 경고", alert_summary(skin.get("sensitiser_alerts"), skin.get("sensitisation_status")) + " / "
                + alert_summary(skin.get("stability_alerts"), skin.get("stability_status")))
            add("구조 기반 안전성 검토", structural.get("overall"))
            add("주요 구조 경고", structural.get("concerns") or ("평가한 구조 규칙에서 경고 기록 없음 · 안전성 미확립" if structural.get("verdicts") else None))
            toxicology = []
            for endpoint in ("Skin_Reaction", "AMES", "Carcinogens_Lagunin", "ClinTox", "DILI", "hERG"):
                values = row.get("tox_" + endpoint)
                rendered = "; ".join(str(item.get("model_id") or "모델") + "=" + brief(item.get("value")) for item in values[:2]) if isinstance(values, list) and values else "미평가"
                toxicology.append(endpoint + ": " + rendered)
            add("독성 모델 출력 · 실측 아님", "\n".join(toxicology))
            additional = mapping(row.get("additional_toxicity"))
            model = mapping(additional.get("model_identity"))
            additional_status = additional.get("status")
            status_labels = {"success": "계산 완료", "partial_nonfinite": "일부 출력 결측",
                "failed": "계산 실패", "failed_invalid_probability": "유효 범위 위반으로 예측 표시 보류",
                "unassessed_model_not_configured": "추가 모델 미설정 · 미평가",
                "unassessed_model_stage_held": "추가 모델 실행 보류 · 미평가",
                "unassessed_model_candidate_limit": "추가 모델 계산 수 제한 · 미평가",
                "not_requested_cached_prediction": "기존 DB 예측이 있어 추가 계산하지 않음",
                "unsupported_structure": "추가 모델이 지원하지 않는 구조"}
            add("추가 로컬 모델 · 별도 채널", status_labels.get(additional_status, "미평가")
                + (" (" + additional_status + ")" if isinstance(additional_status, str) else "")
                + " · 기존 DB 기록과 별도 · 순위 계산 미반영")
            add("추가 모델 식별", " · ".join(key + "=" + brief(model.get(key)) for key in ("provider", "version", "model_id")))
            add("추가 모델 해석 범위", "적용 영역 " + brief(additional.get("applicability_domain", model.get("applicability_domain")), "미확립")
                + " · 실측 독성·피부 효능 근거 아님 · 입체화학 민감도 미확립")
            definitions_sha = additional.get("endpoint_definitions_sha256")
            if not isinstance(definitions_sha, str):
                definitions_sha = None
            if definitions_sha not in additional_endpoint_metadata:
                try:
                    definitions = json.loads(self.service.artifacts.get_bytes(definitions_sha)) if definitions_sha else []
                    additional_endpoint_metadata[definitions_sha] = {entry["id"]: entry for entry in definitions
                        if isinstance(entry, dict) and isinstance(entry.get("id"), str)} if isinstance(definitions, list) else {}
                except (ValueError, TypeError, OSError):
                    additional_endpoint_metadata[definitions_sha] = {}
            definitions = additional_endpoint_metadata[definitions_sha]
            invalid_probabilities = mapping(additional.get("quality_flags")).get("invalid_classifier_probability") or []

            def additional_value(key, value, *, derived):
                if value is None:
                    return "미평가·결측"
                invalid = isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value)
                definition = definitions.get(key, {})
                if not derived:
                    classifier = key in {"Skin_Reaction", "AMES", "Carcinogens_Lagunin", "ClinTox", "DILI", "hERG"} or definition.get("task_type") == "classification"
                    invalid = invalid or key in invalid_probabilities
                    if classifier and not invalid:
                        invalid = not 0 <= value <= 1
                    if not invalid and not classifier and definition.get("task_type") != "regression":
                        return "지표 정의 미확인 · 원자료 참조"
                if invalid:
                    warnings.append({"candidate_id": row["candidate_id"], "component": "additional_toxicity",
                                     "endpoint": key, "reason": "invalid_additional_model_display_value"})
                    return "유효 범위·자료형 위반 · 표시 보류"
                return brief(value) + (" " + str(definition["units"]) if definition.get("units") else "")

            learned = mapping(additional.get("values"))
            if additional_status in {"success", "partial_nonfinite"} and learned:
                items = [key + ": " + additional_value(key, value, derived=False) for key, value in sorted(learned.items())[:200]]
                for offset in range(0, len(items), 6):
                    add("추가 모델 예측값 " + str(offset // 6 + 1), "\n".join(items[offset:offset + 6]))
            else:
                add("추가 모델 예측값", "성공적으로 확인된 추가 예측값 없음 · 실패·보류 원자료는 별도 보존")
            derived = mapping(additional.get("physicochemical_values"))
            items = [key + ": " + additional_value(key, value, derived=True) for key, value in sorted(derived.items())[:200]]
            if items:
                for offset in range(0, len(items), 6):
                    add("추가 물리화학 계산 " + str(offset // 6 + 1), "\n".join(items[offset:offset + 6]) + "\nRDKit 계산값 · 독성 예측과 별도")
            else:
                add("추가 물리화학 계산", "별도 계산값 미기록 · 독성 평가가 아님")
            add("결측·해석 제한", row.get("priority_missing_channels") or "빈 수치·미기록은 0이나 안전 판정이 아닙니다.")
            similarity = mapping(row.get("reference_similarity"))
            comparison = mapping(row.get("reference_alignment")) or mapping(hop.get("comparison")) or hop
            add("호핑 비교", "상태 " + brief(comparison.get("status")) + " · 2D 유사도 " + brief(similarity.get("tanimoto_2d", hop.get("tanimoto_2d")))
                + " · 형태 " + brief(comparison.get("shape_score")) + " · 약리 특징 " + brief(comparison.get("pharmacophore_color_score")))
            novel = row.get("novel_scaffold")
            add("스카폴드 비교 범위", "새 환상 코어: " + ("예" if novel is True else "아니오" if novel is False else "미평가")
                + " · 비고리 연결 구조 변경: " + brief(mapping(row.get("acyclic_topology_comparison")).get("different_from_query_topologies")))
            add("주 도킹 결과", (brief(primary.get("score", primary.get("value"))) + " " + brief(primary.get("unit")) + " · " + brief(primary.get("status"))) if primary else None)
            add("도킹 비교 프로토콜", " · ".join(key + "=" + brief(primary.get(key)) for key in ("receptor_id", "pocket_id", "engine", "protocol_id")))
            add("재도킹 기준 확인", brief(redocking.get("status")) + " · RMSD " + brief(redocking.get("rmsd_angstrom")) + " Å · 기준 " + brief(redocking.get("threshold_angstrom")) + " Å")
            raw_chunk = "chunks/candidates-" + str((count - 1) // 1000 + 1).zfill(5) + ".json"
            add("전체 원자료", raw_chunk + " · 후보 ID " + row["candidate_id"] + " · 모든 원래 필드와 관측값 보존")
            detail.cell(detail.max_row, 2).hyperlink = raw_chunk
            detail.cell(detail.max_row, 2).style = "Hyperlink"
            add("원자료 필드", ", ".join(field for field, _ in FIELDS))
            detail.column_dimensions["A"].width = 32
            detail.column_dimensions["B"].width = 95
            detail.column_dimensions["C"].width = 37
            detail.column_dimensions["D"].width = 38
            detail.freeze_panes = "B4"
            for cells in detail:
                for cell in cells:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                value = str(cells[1].value or "")
                lines = max(sum(max(1, math.ceil(len(part) / 58)) for part in value.splitlines()),
                            sum(max(1, math.ceil(len(part) / 17)) for part in str(cells[2].value or "").splitlines()))
                detail.row_dimensions[cells[0].row].height = min(260, max(30, 18 * lines + 14))
            detail.row_dimensions[1].height = 30
            for cell in detail[1]:
                cell.font = Font(bold=True, color="000000")
            if image_path:
                detail["D3"] = "2D 구조 · 입력 SMILES"
                image = Image(str(image_path))
                image.width, image.height = 240, 140
                detail.add_image(image, "D4")
            else:
                detail["D3"] = row["image_kind"]
            if pose_image_path:
                detail["D12"] = "저장 도킹 포즈 · 좌표 투영"
                image = Image(str(pose_image_path))
                image.width, image.height = 180, 100
                detail.add_image(image, "D13")
            sheet.cell(sheet.max_row, 1).hyperlink = "#'" + title + "'!A1"
            sheet.cell(sheet.max_row, 1).style = "Hyperlink"
            detail_tabs.append({"candidate_id": row["candidate_id"], "workbook": "review-" + str(len(books) + 1).zfill(4) + ".xlsx", "sheet": title})

        def flush_chunk():
            if not chunk:
                return
            name = "chunks/candidates-" + str(len(chunks) + 1).zfill(5) + ".json"
            raw = canonical_json(chunk)
            (directory / name).write_bytes(raw)
            chunks.append(
                {
                    "file": name,
                    "count": len(chunk),
                    "sha256": hashlib.sha256(raw).hexdigest(),
                }
            )
            chunk.clear()

        saved_hashes, expanded_cards, warned_hashes = set(), set(), set()

        def save_hash(sha, *, candidate_id=None, component=None,
                      failure_reason="artifact_quarantined", shared_metadata=False):
            if not isinstance(sha, str) or len(sha) != 64 or any(c not in "0123456789abcdef" for c in sha):
                return None
            path = directory / "assets" / sha
            if sha in saved_hashes and path.exists():
                return path
            saved_hashes.add(sha)
            if not path.exists():
                try:
                    raw = self.service.artifacts.get_bytes(sha)
                    path.write_bytes(raw)
                except (ValueError, OSError):
                    key = (sha, None if shared_metadata else candidate_id, component)
                    if key not in warned_hashes:
                        warned_hashes.add(key)
                        warnings.append({"sha256": sha, "reason": failure_reason,
                            **({"candidate_id": candidate_id} if candidate_id and not shared_metadata else {}),
                            **({"component": component} if component else {})})
                    return None
            return path

        def save_model_metadata(value, key, sha, *, candidate_id):
            """A donor content digest does not imply storage in platform CAS."""
            references = value.get("metadata_refs", {})
            reference = references.get(key) if isinstance(references, dict) else None
            inline_key = {"model_manifest_sha256": "model_manifest",
                          "endpoint_definitions_sha256": "endpoint_definitions"}[key]
            inline = value.get(inline_key)

            def invalid(reason):
                warning_key = (str(sha), None, key, reason)
                if warning_key not in warned_hashes:
                    warned_hashes.add(warning_key)
                    warnings.append({"sha256": sha, "component": key, "reason": reason})

            if not isinstance(sha, str) or len(sha) != 64 or any(c not in "0123456789abcdef" for c in sha):
                invalid("metadata_reference_quarantined")
                return

            if reference is not None:
                if (not isinstance(reference, dict) or reference.get("sha256") != sha
                        or reference.get("storage") not in {"platform_cas", "external_registry"}):
                    invalid("metadata_reference_quarantined")
                    return
                if reference["storage"] == "platform_cas":
                    save_hash(sha, candidate_id=candidate_id, component=key)
                    return
                if (reference.get("source") != "SkinScout model registry"
                        or reference.get("digest_format") not in {
                            "canonical_json", "canonical_json_utf8_v1", "canonical_json_ascii_v1"}):
                    invalid("metadata_reference_quarantined")
                    return
            if (key == "endpoint_definitions_sha256" and reference is None
                    and value.get("source") == "additional_local_cpu_prediction" and isinstance(inline, list)):
                # The additional-model API exposes a bounded display view;
                # its hash identifies the complete definitions in CAS. Accept
                # that view only when every field matches the pinned source.
                try:
                    path = self.service.artifacts.verified_path(sha)
                    if path.stat().st_size > 16 * 1024**2:
                        raise ValueError("Bounded endpoint definitions required")
                    complete = json.loads(path.read_bytes())
                    if not isinstance(complete, list) or len(complete) > 200 or any(not isinstance(entry, dict) for entry in complete):
                        raise ValueError("Typed endpoint definition list required")
                    projection = [{field: entry.get(field) if isinstance(entry.get(field), str)
                                   and len(entry[field]) <= 512 and not any(ord(c) < 32 for c in entry[field]) else None
                                   for field in ("id", "task_type", "category", "name", "units")} for entry in complete]
                    if inline == projection:
                        save_hash(sha, candidate_id=candidate_id, component=key)
                        return
                except (ValueError, TypeError, OSError):
                    pass  # ordinary full-inline validation still quarantines a real conflict
            if inline is not None:
                try:
                    utf8 = canonical_json(inline)
                    ascii_json = json.dumps(inline, sort_keys=True, separators=(",", ":"),
                                            ensure_ascii=True, allow_nan=False).encode()
                    format = reference.get("digest_format") if reference else None
                    allowed = ([ascii_json] if format == "canonical_json_ascii_v1" else [utf8]
                               if format in {"canonical_json", "canonical_json_utf8_v1"} else [utf8, ascii_json])
                    raw = next((data for data in allowed if hashlib.sha256(data).hexdigest() == sha), None)
                    if raw is None:
                        raise ValueError("Inline metadata digest mismatch")
                except (ValueError, TypeError):
                    invalid("inline_metadata_quarantined")
                    return
                (directory / "assets" / sha).write_bytes(raw)
                saved_hashes.add(sha)
            elif reference is None:
                # Historical rows may carry true CAS hashes, so attempt a
                # verified read. Missing storage is uncertainty, not evidence
                # of a quarantined donor model. Warn once per shared digest.
                save_hash(sha, component=key, shared_metadata=True,
                          failure_reason="metadata_storage_unestablished")
            # An explicit external reference remains provenance in candidate
            # JSON; it does not authorize fetching donor files or weights.

        def save_receptor_card(sha, *, candidate_id):
            path = save_hash(sha, candidate_id=candidate_id, component="receptor_identity_card")
            if path is None or sha in expanded_cards:
                return
            expanded_cards.add(sha)
            try:
                if path.stat().st_size > 2 * 1024 * 1024:
                    raise ValueError("bounded receptor identity card required")
                card = json.loads(path.read_bytes())
                entries = card.get("artifacts", {})
                if not isinstance(entries, dict) or len(entries) > 1000:
                    raise ValueError("bounded receptor identity artifact map required")
                # Only exact CAS assets declared present by the frozen native
                # receptor card. Paths never authorize filesystem traversal;
                # catalogs, model weights and global caches are not expanded.
                for name, entry in entries.items():
                    if isinstance(entry, dict) and entry.get("present") is True:
                        save_hash(entry.get("sha256"), candidate_id=candidate_id,
                                  component="receptor_artifact:" + name)
            except (ValueError, TypeError, AttributeError, OSError):
                warnings.append({"sha256": sha, "reason": "receptor_evidence_closure_quarantined",
                                 "candidate_id": candidate_id})

        asset_keys = STRUCTURE_ASSET_KEYS | {
            "diagnostic_log_sha256",
            "source_record_sha256",
            "source_chunk_sha256",
            "target_descriptor_sha256",
            "reference_set_sha256",
            "reference_sdf_sha256",
            "candidate_sdf_sha256",
            "model_manifest_sha256",
            "endpoint_definitions_sha256",
            "identity_card_sha256",
            "uniprot_record_sha256",
            "pose_sha256",
            "input_sha256",
            "receptor_sha256",
            "receptor_identity_card_sha256",
            "preparation_receipt_sha256",
            "ligand_preparation_sha256",
            "pdbqt_sha256",
            "top_pose_sdf_sha256",
            "comparison_reference_set_sha256",
            "reference_observed_sdf_sha256",
            "ref_observed_sdf_sha256",
            "raw_predictions_sha256",
            "exact_input_receipt_sha256",
        }
        expanded_json = set()
        structured_artifacts = {"target_descriptor_sha256", "reference_set_sha256", "comparison_reference_set_sha256",
                                "ligand_preparation_sha256", "literature_artifact_sha256"}

        def save_structured(sha, *, candidate_id, component):
            path = save_hash(sha, candidate_id=candidate_id, component=component)
            if path is None or sha in expanded_json:
                return
            expanded_json.add(sha)
            try:
                if path.stat().st_size > 16 * 1024**2:
                    raise ValueError("Structured evidence expansion exceeds 16 MiB; original blob retained")
                save_evidence(json.loads(path.read_bytes()), candidate_id=candidate_id)
            except (ValueError, TypeError, OSError):
                warnings.append({"sha256": sha, "component": component, "candidate_id": candidate_id,
                                 "reason": "structured_evidence_closure_not_expanded_original_retained"})

        def save_evidence(value, *, candidate_id=None):
            if isinstance(value, dict):
                if value.get("provider") == "admet_ai" and value.get("policy") == "local-admet-ai-exact-input-cpu-v1":
                    return  # pinned runtime/weight file identities; outer receipt owns the CAS metadata
                for key, item in value.items():
                    if key in {"metadata_refs", "donor_metadata_refs"}:
                        continue  # storage descriptors are read with their owning hashes
                    # The first canary used this name for a semantic reference
                    # policy hash as well as the actual CAS artifact hash. Its
                    # verified snapshot distinguishes the two; retain the
                    # semantic ID in JSON without fabricating a missing file.
                    provenance = source.get("provenance", {})
                    if (
                        key == "reference_set_sha256"
                        and item == provenance.get("reference_policy_sha256")
                        and item != provenance.get("reference_set_sha256")
                    ):
                        continue
                    if key in {"identity_card_sha256", "receptor_identity_card_sha256"}:
                        save_receptor_card(item, candidate_id=candidate_id)
                    elif key in {"model_manifest_sha256", "endpoint_definitions_sha256"}:
                        save_model_metadata(value, key, item, candidate_id=candidate_id)
                    elif key == "sdf_artifact" and isinstance(item, dict):
                        save_hash(item.get("sha256"), candidate_id=candidate_id, component="saved_pose_sdf")
                    elif key in structured_artifacts:
                        save_structured(item, candidate_id=candidate_id, component=key)
                    elif key in asset_keys:
                        save_hash(item, candidate_id=candidate_id, component=key)
                    elif isinstance(item, (dict, list)):
                        save_evidence(item, candidate_id=candidate_id)
            elif isinstance(value, list):
                for item in value:
                    save_evidence(item, candidate_id=candidate_id)

        save_evidence(source.get("provenance", {}))
        for row in self.service.catalog.iter_candidates(
            request["snapshot_id"], target_id=target
        ):
            stop()
            if count % workbook_limit == 0:
                flush_book()
                workbook, sheet = new_book()
            cid = row["candidate_id"]
            save_evidence(row, candidate_id=cid)
            for prediction in row.get("admet") or []:
                if not isinstance(prediction, dict):
                    continue
                model = prediction.get("model_id")
                sha = prediction.get("endpoint_definitions_sha256")
                if model and model not in endpoint_metadata:
                    definitions = prediction.get("endpoint_definitions")
                    if definitions is None and sha:
                        try:
                            definitions = json.loads(
                                self.service.artifacts.get_bytes(sha)
                            )
                        except ValueError:
                            definitions = None
                    if isinstance(definitions, list):
                        endpoint_metadata[model] = definitions
            ids.append(cid)
            count += 1
            pose = None
            observations = row.get("docking") or []
            observations = (
                [observations] if isinstance(observations, dict) else observations
            )
            for observation in observations:
                for key in (
                    "pose_sha256",
                    "input_sha256",
                    "receptor_sha256",
                    "receptor_identity_card_sha256",
                ):
                    save_hash(observation.get(key))
            primary = row.get("primary_docking") or {}
            index = primary.get("observation_index", 0)
            if observations:
                try:
                    pose_view = candidate_docking_view(
                        self.service.catalog,
                        self.service.artifacts,
                        cid,
                        snapshot_id=request["snapshot_id"],
                        target_id=target,
                        observation_index=index,
                    )
                    name = "views/" + hashlib.sha256(cid.encode()).hexdigest() + ".json"
                    (directory / name).write_bytes(canonical_json(pose_view))
                    row["offline_pose_view"] = name
                    pose = pose_view.get("ligand")
                except ValueError as error:
                    row["offline_pose_error"] = str(error)
                    warnings.append(
                        {"candidate_id": cid, "reason": "pose_view_unavailable"}
                    )
            image_path, pose_image_path = None, None
            smiles = row.get("smiles")
            molecule = Chem.MolFromSmiles(smiles) if isinstance(smiles, str) and smiles.strip() else None
            if molecule is not None and molecule.GetNumAtoms() == 0:
                molecule = None
            row["image"] = None
            row["image_kind"] = "2D 구조 없음 · SMILES 미제공" if not smiles else "2D 구조 없음 · SMILES 해석 불가"
            if molecule is not None:
                try:
                    depiction = Chem.Mol(molecule)
                    rdDepictor.Compute2DCoords(depiction, clearConfs=True, forceRDKit=True)
                    buffer = io.BytesIO()
                    Draw.MolToImage(depiction, size=(480, 280)).save(buffer, format="PNG")
                    relative = "images/" + hashlib.sha256(cid.encode()).hexdigest() + "-2d.png"
                    image_path = directory / relative
                    image_path.write_bytes(buffer.getvalue())
                    row["image"] = relative
                    row["image_kind"] = "2D 구조 · 입력 SMILES"
                    add_structure_image(sheet, image_path, row=sheet.max_row + 1, column=2)
                except (ValueError, RuntimeError):
                    row["image_kind"] = "2D 구조 생성 실패 · 원래 SMILES 참조"
                    warnings.append({"candidate_id": cid, "reason": "2D_depiction_failed"})
            else:
                warnings.append({"candidate_id": cid, "reason": "2D_structure_unavailable"})
            row["pose_image"] = None
            row["pose_image_kind"] = None
            if pose and pose.get("data"):
                # Reuse the audited coordinate-projection renderer, without a
                # generated conformer fallback (undocked molecules use 2D).
                try:
                    import importlib.util

                    path = export_asset_root() / "bridge/export_images.py"
                    spec = importlib.util.spec_from_file_location(
                        "_skinpipe_export_image", path
                    )
                    module = importlib.util.module_from_spec(spec)
                    spec.loader.exec_module(module)
                    png, metadata = module.thumbnail(None, pose)
                    if png:
                        relative = "images/" + hashlib.sha256(cid.encode()).hexdigest() + "-pose.png"
                        pose_image_path = directory / relative
                        pose_image_path.write_bytes(png)
                        row["pose_image"] = relative
                        row["pose_image_kind"] = "저장 도킹 포즈 · 좌표 투영"
                except (ValueError, OSError):
                    pass
            row.update(toxicity_columns(row, warnings))
            row.update(review_columns(row))
            values = []
            for field, _ in FIELDS:
                values.append(xlsx_value(row.get(field), candidate_id=cid, field=field, warnings=warnings))
            sheet.append(values)
            sheet.row_dimensions[sheet.max_row].height = 114
            for cell in sheet[sheet.max_row]:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            for i, (field, _) in enumerate(FIELDS, 1):
                if field in {"priority_score_v2", "review_skin_score", "review_logp", "review_tpsa", "review_mw"}:
                    sheet.cell(sheet.max_row, i).number_format = "0.00"
                elif field.startswith("review_tox_") or field in {"review_solubility", "review_qed"}:
                    sheet.cell(sheet.max_row, i).number_format = "0.0000"
            if with_details:
                candidate_sheet(row, image_path, pose_image_path)
            # SDF order matches JSON/XLSX. Structurally invalid rows get a
            # separate explicit ledger rather than a fabricated molecule.
            if molecule is not None:
                coordinate_scope = "source exact structure; 2D"
                if pose and pose.get("format") == "sdf" and pose.get("data"):
                    saved = Chem.MolFromMolBlock(
                        pose["data"].split("$$$$")[0], removeHs=False
                    )
                    if saved is not None:
                        molecule = saved
                        coordinate_scope = "stored prepared docking pose; prepared chemistry may differ from source material"
                molecule.SetProp("candidate_id", cid)
                molecule.SetProp("target_id", target)
                sdf.write(
                    Chem.MolToMolBlock(molecule)
                    + "\n> <candidate_id>\n"
                    + cid
                    + "\n\n> <target_id>\n"
                    + target
                    + "\n\n> <coordinate_scope>\n"
                    + coordinate_scope
                    + "\n\n$$$$\n"
                )
            else:
                warnings.append(
                    {"candidate_id": cid, "reason": "SDF_structure_unavailable"}
                )
            chunk.append(row)
            if len(chunk) == 1000:
                flush_chunk()
            progress(
                {
                    "stage": "export",
                    "processed": count,
                    "total": source["target_counts"][target],
                }
            )
        flush_chunk()
        if workbook is None:
            workbook, sheet = new_book()
        flush_book()
        if count != source["target_counts"].get(target, 0) or len(ids) != len(set(ids)):
            raise ValueError(
                "Export population count or IDs differ from source snapshot"
            )
        sdf.flush()
        (directory / "source-manifest.json").write_bytes(canonical_json(source))
        public = export_asset_root() / "public"
        for name in ("run-offline.html", "run-offline.js", "run-offline-worker.js"):
            shutil.copyfile(
                public / name,
                directory / ("index.html" if name.endswith(".html") else name),
            )
        vendor = public / "vendor/molstar"
        if vendor.exists():
            shutil.copytree(vendor, directory / "vendor/molstar")
        else:
            raise ValueError(
                "Pinned local MolStar assets missing; offline export deferred"
            )
        receipt = {
            "schema": "skinpipe-run-export-v1",
            "exporter": request.get("exporter", self.implementation_identity()),
            "snapshot_id": request["snapshot_id"],
            "target_id": target,
            "source_manifest_sha256": digest(source),
            "candidate_count": count,
            "publication_scope": source.get("provenance", {}).get("publication_scope", {
                "status": "NOT_RECORDED", "whole_workflow_complete": None}),
            "container_format": "zip",
            "requested_format": request.get("format", "zip"),
            "retry_history": request.get("retry_history"),
            "filters": source.get("provenance", {}).get("filters", {}),
            "admet_search_policy": source.get("provenance", {}).get(
                "admet_search_policy"
            ),
            "candidate_ids_in_order": ids,
            "candidate_order_sha256": digest(ids),
            "chunks": chunks,
            "workbooks": books,
            "workbook_row_limit": workbook_limit,
            "candidate_detail_tabs": detail_tabs,
            "research_input": provenance.get("research_input", provenance.get("input", provenance.get("target_descriptor"))),
            "skin_context": provenance.get("skin_context"),
            "catalog_selection": provenance.get("catalog_selection"),
            "warnings": warnings,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "scope": "entire saved target population including filter exclusions; shared global catalog is separate",
            "external_network": "blocked",
            "raw_missing_values": "null",
            "scientific_validation": "NOT_ESTABLISHED",
            "files": {
                str(p.relative_to(directory)): {
                    "sha256": sha256_file(p),
                    "size_bytes": p.stat().st_size,
                }
                for p in directory.rglob("*")
                if p.is_file()
            },
        }
        (directory / "manifest.json").write_bytes(canonical_json(receipt))
        return receipt
