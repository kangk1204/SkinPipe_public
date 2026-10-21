"""Submission contention regressions against an explicitly isolated PostgreSQL DB."""
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from http.client import HTTPConnection
import json
import os
from pathlib import Path
import tempfile
import threading
import time
from types import SimpleNamespace
import unittest
import uuid

import psycopg
from psycopg.conninfo import make_conninfo
from psycopg.rows import dict_row

from skinpipe.platform.queue import Queue, SCHEMA, SubmissionBusy
from skinpipe.platform.personal_local import PersonalServer


class SubmissionLockTests(unittest.TestCase):
    def setUp(self):
        dsn = os.environ.get("SKINPIPE_RUNTIME_AUDIT_DSN") or os.environ.get("SKINPIPE_TEST_PG_DSN")
        config = os.environ.get("SKINPIPE_TEST_POSTGRES_CONFIG")
        if not dsn and config:
            dsn = make_conninfo(**json.loads(Path(config).read_text()))
        if not dsn:
            self.skipTest("Explicit isolated PostgreSQL test configuration required")
        self.queue = Queue(dsn, initialize=False)
        with self.queue.connection() as conn:
            self.assertTrue(conn.execute("SELECT current_database() AS name").fetchone()["name"].endswith("_test"))
            conn.execute(SCHEMA)
        self.owner = "lock-test-" + uuid.uuid4().hex

    def submit(self, queue=None, **changes):
        values = dict(owner=self.owner, project=self.owner, idempotency_key="same", budget_seconds=100)
        values.update(changes)
        return (queue or self.queue).enqueue("probe", {"value": 1}, **values)

    def hold_key(self, conn):
        conn.execute("SELECT pg_advisory_xact_lock(hashtextextended(%s,0))", (self.owner + ":same",))

    def fast_timeout_queue(self):
        parent = self.queue

        class ShortWaitQueue(Queue):
            @contextmanager
            def connection(self):
                with parent.connection() as conn:
                    conn.execute("SET LOCAL lock_timeout = '200ms'")
                    yield conn

        return ShortWaitQueue(parent.dsn, initialize=False)

    def test_committed_replay_does_not_wait_or_reset_terminal_job(self):
        original = self.submit()
        with self.queue.connection() as conn:
            conn.execute("UPDATE platform_jobs SET status='complete', attempts=1, compute_seconds=7, lease_token='private-test-token' WHERE id=%s", (original["id"],))
            conn.execute("UPDATE platform_compute_budgets SET used_seconds=7 WHERE id=%s", (original["budget_scope"],))
        with self.queue.connection() as holder:
            self.hold_key(holder)
            replay = self.submit(self.fast_timeout_queue())
            self.assertEqual(replay["id"], original["id"])
            self.assertEqual((replay["status"], replay["attempts"], replay["compute_seconds"]), ("complete", 1, 7))
            self.assertNotIn("lease_token", replay)
            with self.assertRaisesRegex(ValueError, "different request"):
                self.submit(self.fast_timeout_queue(), budget_seconds=101)
        with self.queue.connection() as conn:
            self.assertEqual(conn.execute("SELECT used_seconds FROM platform_compute_budgets WHERE id=%s", (original["budget_scope"],)).fetchone()["used_seconds"], 7)
            self.assertEqual(conn.execute("SELECT count(*) AS n FROM platform_job_events WHERE job_id=%s", (original["id"],)).fetchone()["n"], 1)

    def test_concurrent_first_submissions_publish_one_job_and_event(self):
        with ThreadPoolExecutor(6) as pool:
            jobs = list(pool.map(lambda _: self.submit(), range(6)))
        self.assertEqual(len({job["id"] for job in jobs}), 1)
        with self.queue.connection() as conn:
            self.assertEqual(conn.execute("SELECT count(*) AS n FROM platform_job_events WHERE job_id=%s", (jobs[0]["id"],)).fetchone()["n"], 1)

    def test_waiter_reads_creator_commit_after_acquiring_lock(self):
        with ThreadPoolExecutor(1) as pool:
            with self.queue.connection() as creator:
                self.hold_key(creator)

                class CreatingQueue(Queue):
                    @contextmanager
                    def connection(self):
                        yield creator

                original = self.submit(CreatingQueue(self.queue.dsn, initialize=False))
                future = pool.submit(self.submit)
                # Observe the actual blocked backend before releasing the creator.
                with psycopg.connect(self.queue.dsn, autocommit=True, row_factory=dict_row) as monitor:
                    deadline = time.monotonic() + 3
                    while time.monotonic() < deadline:
                        if monitor.execute("SELECT 1 FROM pg_stat_activity WHERE %s=ANY(pg_blocking_pids(pid))", (creator.info.backend_pid,)).fetchone():
                            break
                        time.sleep(0.01)
                    else:
                        self.fail("Concurrent submitter never waited for the creator")
            self.assertEqual(future.result(timeout=3)["id"], original["id"])

    def test_new_submission_timeout_can_be_retried_with_same_key(self):
        with self.queue.connection() as holder:
            self.hold_key(holder)
            with self.assertRaises(SubmissionBusy):
                self.submit(self.fast_timeout_queue())
        self.assertEqual(self.queue.list(owner=self.owner), [])
        original = self.submit()
        self.assertEqual(self.submit()["id"], original["id"])

    def test_late_budget_lock_rolls_back_user_project_and_job(self):
        with self.queue.connection() as holder:
            holder.execute("INSERT INTO platform_compute_budgets(id,owner_id,limit_seconds) VALUES(%s,%s,100)", (self.owner + ":same", self.owner))
            with self.assertRaises(SubmissionBusy):
                self.submit(self.fast_timeout_queue())
            with self.queue.connection() as conn:
                for table in ("platform_users", "platform_projects"):
                    # Table names are fixed test constants, never user input.
                    self.assertIsNone(conn.execute(f"SELECT id FROM {table} WHERE id=%s", (self.owner,)).fetchone())
                self.assertEqual(self.queue.list(owner=self.owner), [])
        self.assertEqual(self.submit()["status"], "queued")

    def test_completed_excel_http_download_survives_held_submission_lock(self):
        # Exercise the real HTTP handler, queue, exporter, CAS and workbook.
        # Only the two candidate records are synthetic; no running study is used.
        import io
        import zipfile
        from openpyxl import load_workbook
        from psycopg import sql
        from skinpipe.platform.artifacts import ArtifactStore
        from skinpipe.platform.catalog import Catalog
        from skinpipe.platform.exports import RunExports

        database = "export_http_" + uuid.uuid4().hex + "_test"
        with psycopg.connect(self.queue.dsn, autocommit=True) as admin:
            admin.execute(sql.SQL("CREATE DATABASE {}").format(sql.Identifier(database)))
        try:
            queue = Queue(make_conninfo(self.queue.dsn, dbname=database))
            with tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                catalog = Catalog(root / "catalog")
                manifest = catalog.create_snapshot([
                    {"candidate_id": "synthetic-kept", "target_id": "compound:synthetic", "smiles": "CCO", "discovery_filter_pass": True},
                    {"candidate_id": "synthetic-excluded", "target_id": "compound:synthetic", "smiles": "CCN", "discovery_filter_pass": False},
                ], provenance={"owner_id": "researcher"})
                service = SimpleNamespace(root=root, catalog=catalog, artifacts=ArtifactStore(root / "cas"), queue=queue)
                service.exports = RunExports(service)
                service.post = lambda path, body, owner: service.exports.submit(body, owner=owner)
                public = root / "public"
                public.mkdir()
                (public / "local-research.html").write_text("<title>Synthetic export regression</title>")
                server = PersonalServer(service, {"discovery": {}}, port=0, asset_root=root)
                thread = threading.Thread(target=server.serve_forever, daemon=True)
                thread.start()
                conn = HTTPConnection("127.0.0.1", server.server_port, timeout=10)
                try:
                    conn.request("GET", "/")
                    response = conn.getresponse()
                    cookie = response.getheader("Set-Cookie").split(";", 1)[0]
                    response.read()
                    headers = {"Cookie": cookie, "Content-Type": "application/json", "X-SkinPipe-Local": "1"}
                    body = json.dumps({"snapshot_id": manifest["snapshot_id"], "target_id": "compound:synthetic", "format": "xlsx", "candidate_details": True})

                    def submit_http():
                        conn.request("POST", "/v2/exports", body, headers)
                        response = conn.getresponse()
                        value = json.loads(response.read())
                        self.assertEqual(response.status, 200, value)
                        return value

                    original = submit_http()
                    job = queue.claim("synthetic-export-worker", ["result_export"])
                    self.assertEqual(job["id"], original["export_id"])
                    result = service.exports.execute(job["request"], stop=lambda: None, progress=lambda value: None)
                    queue.finish(job["id"], job["lease_token"], result, compute_seconds=1)
                    frozen = queue.get(job["id"])
                    with queue.connection() as holder:
                        holder.execute("SELECT pg_advisory_xact_lock(hashtextextended(%s,0))", ("researcher:" + job["idempotency_key"],))
                        replay = submit_http()
                        self.assertEqual((replay["export_id"], replay["status"]), (job["id"], "complete"))
                        conn.request("GET", "/v2/exports/" + job["id"] + "/download", headers={"Cookie": cookie})
                        response = conn.getresponse()
                        downloaded = response.read()
                        self.assertEqual(response.status, 200)
                        self.assertEqual(response.getheader("Content-Type"), "application/zip")
                    self.assertEqual(queue.get(job["id"]), frozen)
                    self.assertEqual(downloaded, service.artifacts.verified_path(result["sha256"]).read_bytes())
                    with zipfile.ZipFile(io.BytesIO(downloaded)) as bundle:
                        self.assertIsNone(bundle.testzip())
                        book = load_workbook(io.BytesIO(bundle.read("review-0001.xlsx")))
                        try:
                            self.assertEqual(book["후보 검토"].max_row, 3)
                            self.assertEqual([name for name in book.sheetnames if name.startswith("C000")], ["C000001", "C000002"])
                            self.assertEqual({book[name]["B1"].value for name in ("C000001", "C000002")}, {"synthetic-kept", "synthetic-excluded"})
                        finally:
                            book.close()
                finally:
                    conn.close()
                    server.shutdown()
                    server.server_close()
                    thread.join(timeout=5)
        finally:
            with psycopg.connect(self.queue.dsn, autocommit=True) as admin:
                admin.execute(sql.SQL("DROP DATABASE {}").format(sql.Identifier(database)))

    def test_replay_still_checks_project_and_budget_ownership(self):
        job = self.submit()
        with self.queue.connection() as conn:
            conn.execute("INSERT INTO platform_users(id) VALUES(%s)", (self.owner + "-foreign",))
            conn.execute("UPDATE platform_projects SET owner_id=%s WHERE id=%s", (self.owner + "-foreign", self.owner))
        with self.assertRaisesRegex(ValueError, "Project belongs"):
            self.submit()
        with self.queue.connection() as conn:
            conn.execute("UPDATE platform_projects SET owner_id=%s WHERE id=%s", (self.owner, self.owner))
            conn.execute("UPDATE platform_compute_budgets SET owner_id=%s WHERE id=%s", (self.owner + "-foreign", job["budget_scope"]))
        with self.assertRaisesRegex(ValueError, "Compute budget belongs"):
            self.submit()


class SubmissionBusyHTTPTests(unittest.TestCase):
    def test_browser_receives_retryable_response_without_database_details(self):
        def busy(*args, **kwargs):
            raise SubmissionBusy("private database diagnostic must not be exposed")

        with tempfile.TemporaryDirectory() as directory:
            server = PersonalServer(SimpleNamespace(post=busy), {"discovery": {}}, port=0, asset_root=Path(directory))
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            conn = HTTPConnection("127.0.0.1", server.server_port, timeout=5)
            try:
                conn.request("POST", "/v2/exports", "{}", {
                    "Content-Type": "application/json", "X-SkinPipe-Local": "1",
                    "Cookie": "skinpipe_local=" + server.session,
                })
                response = conn.getresponse()
                value = json.loads(response.read())
                self.assertEqual(response.status, 503)
                self.assertEqual(value["error_type"], "SubmissionBusy")
                self.assertTrue(value["retryable"])
                self.assertIn("같은 버튼", value["error"])
                self.assertNotIn("private database", value["error"])
            finally:
                conn.close()
                server.shutdown()
                server.server_close()
                thread.join(timeout=5)
