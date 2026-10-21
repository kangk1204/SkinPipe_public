"""Exercise actual Excel exports with synthetic structures and stored coordinates."""
# ruff: noqa: E402 -- Resolve the checkout before importing its exporter.
import hashlib
import json
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "skinpipe" if (ROOT / "skinpipe/skinpipe").is_dir() else ROOT
sys.path.insert(0, str(SOURCE))

from openpyxl import load_workbook
from rdkit import Chem
from rdkit.Geometry import Point3D

from skinpipe.platform.artifacts import ArtifactStore
from skinpipe.platform.catalog import Catalog
from skinpipe.platform.exports import RunExports, METRIC_GUIDE


class ExcelStructureTests(unittest.TestCase):
    def test_singleton_prediction_and_score_scope_survive_actual_workbook_export(self):
        row = self.row("synthetic", "CCO")
        row.update(role="query_seed", priority_score_v2=60.269, priority_score_observed_channels=84.23,
                   priority_imputed_weight=0.7, priority_missing_channels=["measured", "structural"],
                   discovery_terms={"topical": 50}, skin_goal_assessment={"permeability_ranking": "neutral"},
                   offers=[{"product_url": "https://example.org/compound/synthetic", "supplier": "Synthetic"}],
                   admet={"model_id": "admet-ai-synthetic", "status": "success", "values": {"AMES": 0.2}})
        with tempfile.TemporaryDirectory() as temporary:
            output, result, saved = self.export(Path(temporary), [row])
            self.assertEqual(saved[0]["admet"], row["admet"])
            self.assertEqual(saved[0]["priority_score_v2"], 60.269)
            book = load_workbook(output / result["workbooks"][0])
            try:
                table = book["후보 검토"]
                columns = {cell.value.split(" · ")[0]: table.cell(2, cell.column).value for cell in table[1]}
                for title, expected in (("역할 구분", "출발구조"), ("결측 중립값 비중 (%)", 70),
                                        ("관측 채널만의 점수 (0–100점)", 84.23)):
                    self.assertEqual(columns[title], expected)
                self.assertEqual(next(value for title, value in columns.items() if title.startswith("AMES 예측")), 0.2)
                self.assertIn("중립 50", columns["탐색 점수에 반영한 피부 항목"])
                self.assertIn("neutral_imputation", columns["우선순위 계산 범위"])
                purchase = next(cell for cell in table[2] if cell.value == "https://example.org/compound/synthetic")
                self.assertEqual(purchase.hyperlink.target, purchase.value)
                links = [cell.hyperlink.target for cells in book["C000001"] for cell in cells if cell.hyperlink]
                self.assertIn("https://example.org/compound/synthetic", links)
            finally:
                book.close()

    def export(self, root, rows, *, details=True):
        catalog = Catalog(root / "catalog")
        store = ArtifactStore(root / "artifacts")
        manifest = catalog.create_snapshot(rows)
        frozen = list(catalog.iter_candidates(manifest["snapshot_id"], target_id="fixture"))
        exporter = RunExports(SimpleNamespace(root=root, catalog=catalog, artifacts=store))
        output = root / "export"
        result = exporter.build({"snapshot_id": manifest["snapshot_id"], "target_id": "fixture",
                                 "format": "xlsx", "candidate_details": details}, output)
        self.assertEqual(frozen, list(catalog.iter_candidates(manifest["snapshot_id"], target_id="fixture")))
        saved = []
        for chunk in result["chunks"]:
            saved.extend(json.loads((output / chunk["file"]).read_bytes()))
        return output, result, saved

    @staticmethod
    def row(identifier, smiles):
        return {"candidate_id": identifier, "target_id": "fixture", "smiles": smiles,
                "name": "Synthetic structure", "priority_score_v2": 73.125,
                "admet": [], "docking": [], "discovery_filter_pass": True}

    def test_docked_compound_has_smiles_2d_and_separate_original_pose(self):
        smiles = "C[C@H](O)C(=O)[O-].[Na+]"
        molecule = Chem.MolFromSmiles(smiles)
        conformer = Chem.Conformer(molecule.GetNumAtoms())
        for i in range(molecule.GetNumAtoms()):
            conformer.SetAtomPosition(i, Point3D(i * 0.8, (i % 3) * 0.4, (i % 2) * 0.7))
        conformer.Set3D(True)
        molecule.AddConformer(conformer)
        block = Chem.MolToMolBlock(molecule)
        pose = {"format": "sdf", "data": block, "source_sha256": hashlib.sha256(block.encode()).hexdigest()}
        row = self.row("docked", smiles)
        row["docking"] = [{"engine": "synthetic", "affinity": -6.125, "status": "complete"}]
        with tempfile.TemporaryDirectory() as temporary, patch(
            "skinpipe.platform.pose_view.candidate_docking_view", return_value={"ligand": pose}
        ):
            output, result, saved = self.export(Path(temporary), [row])
            book = load_workbook(output / result["workbooks"][0])
            try:
                table, detail = book["후보 검토"], book["C000001"]
                self.assertEqual([c.value for c in table[1]][2:4], ["2D 구조", "SMILES"])
                self.assertEqual(table["D2"].value, smiles)
                self.assertEqual(len(table._images), 1)
                self.assertEqual(table._images[0].anchor._from.col, 2)
                self.assertEqual(table._images[0].anchor._from.row, 1)
                self.assertEqual(len(detail._images), 2)
                self.assertEqual(detail["D3"].value, "2D 구조 · 확인·저장한 SMILES")
                self.assertEqual(detail["D12"].value, "저장 도킹 포즈 · 좌표 투영")
                fields = {r[0].value: r[1].value for r in detail}
                self.assertEqual(fields["SMILES · 원래 입력 구조"], smiles)
                self.assertEqual(saved[0]["priority_score_v2"], row["priority_score_v2"])
                self.assertEqual(saved[0]["docking"], row["docking"])
                self.assertEqual(saved[0]["smiles"], smiles)
                self.assertNotEqual((output / saved[0]["image"]).read_bytes(),
                                    (output / saved[0]["pose_image"]).read_bytes())
                from PIL import Image
                with Image.open(output / saved[0]["image"]) as depiction:
                    self.assertEqual(depiction.size, (480, 280))
                exported = Chem.MolFromMolBlock((output / "candidates.sdf").read_text().split("$$$$")[0], removeHs=False)
                self.assertEqual(Chem.MolToSmiles(exported), Chem.MolToSmiles(molecule))
                self.assertTrue(exported.GetConformer().Is3D())
                for i in range(molecule.GetNumAtoms()):
                    self.assertLess((exported.GetConformer().GetAtomPosition(i)
                                     - molecule.GetConformer().GetAtomPosition(i)).Length(), 0.0001)
            finally:
                book.close()

    def test_invalid_and_missing_smiles_keep_rows_without_fabricated_images(self):
        rows = [self.row("valid", "c1ccccc1O"), self.row("invalid", "not-a-smiles"), self.row("missing", None)]
        rows[0]["discovery_filter_pass"] = False
        with tempfile.TemporaryDirectory() as temporary:
            output, result, saved = self.export(Path(temporary), rows)
            self.assertEqual(result["candidate_count"], 3)
            self.assertEqual(len(result["candidate_detail_tabs"]), 3)
            by_id = {r["candidate_id"]: r for r in saved}
            self.assertEqual({r["candidate_id"]: r["smiles"] for r in saved},
                             {r["candidate_id"]: r["smiles"] for r in rows})
            self.assertFalse(by_id["valid"]["discovery_filter_pass"])
            self.assertTrue(by_id["valid"]["image"])
            self.assertIsNone(by_id["invalid"]["image"])
            self.assertIsNone(by_id["missing"]["image"])
            self.assertIn("해석 불가", by_id["invalid"]["image_kind"])
            self.assertIn("미제공", by_id["missing"]["image_kind"])
            book = load_workbook(output / result["workbooks"][0])
            try:
                self.assertEqual(len(book["후보 검토"]._images), 1)
                for entry in result["candidate_detail_tabs"]:
                    self.assertEqual(len(book[entry["sheet"]]._images), int(entry["candidate_id"] == "valid"))
            finally:
                book.close()

    def test_split_workbooks_keep_each_image_with_its_smiles(self):
        rows = [self.row(str(i), smiles) for i, smiles in enumerate(["CCO", "CCN", "CC(=O)O", "c1ccccc1", "C[C@@H](O)C(=O)O"])]
        with tempfile.TemporaryDirectory() as temporary, patch("skinpipe.platform.exports.DETAILS_PER_WORKBOOK", 2):
            output, result, saved = self.export(Path(temporary), rows)
            self.assertEqual(len(result["workbooks"]), 3)
            seen = []
            for name in result["workbooks"]:
                book = load_workbook(output / name)
                try:
                    table = book["후보 검토"]
                    self.assertEqual(len(table._images), table.max_row - 1)
                    for number in range(2, table.max_row + 1):
                        candidate = next(r for r in saved if r["candidate_id"] == table.cell(number, 1).value)
                        self.assertEqual(table.cell(number, 4).value, candidate["smiles"])
                        image = next(i for i in table._images if i.anchor._from.row == number - 1)
                        self.assertEqual(hashlib.sha256(image._data()).hexdigest(),
                                         hashlib.sha256((output / candidate["image"]).read_bytes()).hexdigest())
                        seen.append(candidate["candidate_id"])
                finally:
                    book.close()
            self.assertEqual(seen, [r["candidate_id"] for r in rows])

    def test_table_only_export_also_has_2d_before_smiles(self):
        with tempfile.TemporaryDirectory() as temporary:
            output, result, _ = self.export(Path(temporary), [self.row("plain", "CCO")], details=False)
            self.assertEqual(result["candidate_detail_tabs"], [])
            book = load_workbook(output / result["workbooks"][0])
            try:
                self.assertEqual(book["후보 검토"]["C1"].value, "2D 구조")
                self.assertEqual(book["후보 검토"]["D2"].value, "CCO")
                self.assertEqual(len(book["후보 검토"]._images), 1)
            finally:
                book.close()

    def test_calculation_columns_keep_zero_missing_invalid_and_raw_values_distinct(self):
        row = self.row("metrics", "CCO")
        row["skin"] = {"topical_score": 0, "permeation_class": "moderate", "sensitisation_status": "EVALUATED",
                       "sensitiser_alerts": [], "stability_status": "EVALUATED", "stability_alerts": ["hydrolysis:ester"]}
        row["admet"] = [{"model_id": "synthetic-a", "status": "success", "values": {
            "Skin_Reaction": 0, "AMES": 0.123456789, "hERG": 2, "molecular_weight": 46.069,
            "logP": -0.0014, "tpsa": 20.23, "hydrogen_bond_donors": 1, "hydrogen_bond_acceptors": 1}}]
        row["additional_toxicity"] = {"status": "success", "values": {"DILI": 0.9}}
        with tempfile.TemporaryDirectory() as temporary:
            output, result, saved = self.export(Path(temporary), [row])
            book = load_workbook(output / result["workbooks"][0])
            try:
                table = book["후보 검토"]
                cells = {table.cell(1, i).value.splitlines()[0]: table.cell(2, i) for i in range(1, table.max_column + 1)}
                self.assertEqual(cells["피부 적용 점수 (0–10점)"].value, 0)
                self.assertEqual(cells["피부 감작성 예측 (0–1)"].value, 0)
                self.assertEqual(cells["AMES 예측 (0–1)"].value, 0.123456789)
                self.assertEqual(cells["피부 감작성 예측 (0–1)"].data_type, "n")
                self.assertEqual(cells["hERG 예측 (0–1)"].value, "값 오류 · 원자료 확인")
                self.assertEqual(cells["간 손상 예측 (0–1)"].value, "미평가")
                self.assertEqual(cells["분자량 (g/mol)"].value, 46.069)
                self.assertEqual(cells["logP"].value, -0.0014)
                self.assertEqual(cells["감작성 구조 경보"].value, "경보 없음")
                self.assertEqual(saved[0]["admet"], row["admet"])
                self.assertEqual(saved[0]["additional_toxicity"], row["additional_toxicity"])
            finally:
                book.close()

    def test_multiple_models_are_named_instead_of_averaged(self):
        row = self.row("models", "CCO")
        row["admet"] = [{"model_id": "synthetic-a", "status": "success", "values": {"AMES": 0.1}},
                        {"model_id": "synthetic-b", "status": "success", "values": {"AMES": 0.7}},
                        {"model_id": "failed-model", "status": "failed", "values": {"AMES": 0.9}}]
        with tempfile.TemporaryDirectory() as temporary:
            output, result, saved = self.export(Path(temporary), [row])
            book = load_workbook(output / result["workbooks"][0])
            try:
                table = book["후보 검토"]
                columns = {table.cell(1, i).value.splitlines()[0]: i for i in range(1, table.max_column + 1)}
                self.assertEqual(table.cell(2, columns["AMES 예측 (0–1)"]).value,
                                 "synthetic-a: 0.1\nsynthetic-b: 0.7")
                self.assertIn("failed-model: 실패", table.cell(2, columns["ADMET 계산 상태"]).value)
                self.assertEqual(saved[0]["admet"], row["admet"])
            finally:
                book.close()

    def test_metric_directions_distinguish_scores_ranks_and_properties(self):
        import re
        browser = (SOURCE.parent / "web/public/column-help.js").read_text()
        definitions = json.loads(re.search(r"export const metricColumnHelp = (\{.*?\});", browser, re.S).group(1))
        for field, _label, units, direction, meaning, comparison, _source in METRIC_GUIDE:
            self.assertIn(units, definitions[field][0])
            self.assertEqual(meaning, definitions[field][1])
            self.assertEqual(direction + ". " + comparison, definitions[field][2])
        with tempfile.TemporaryDirectory() as temporary:
            output, result, _ = self.export(Path(temporary), [self.row("directions", "CCO")])
            book = load_workbook(output / result["workbooks"][0])
            try:
                table, guide = book["후보 검토"], book["지표 읽는 법"]
                self.assertEqual(book.sheetnames[1], "지표 읽는 법")
                columns = {c.value.splitlines()[0]: c for c in table[1]}
                self.assertIn("↑ 높을수록", columns["평가 점수 (0–100점)"].value)
                self.assertIn("↓ 더 음수", columns["도킹 (kcal/mol)"].value)
                self.assertIn("↓ 낮을수록", columns["AMES 예측 (0–1)"].value)
                self.assertIn("적정 범위", columns["logP"].value)
                self.assertIn("목표 부위", columns["피부 투과 분류"].value)
                self.assertEqual(len(list(guide.iter_rows(min_row=7))), len(METRIC_GUIDE))
                self.assertIn("↓ 1위", guide["C8"].value)
                for field, label in [("priority_score_v2", "평가 점수 (0–100점)"), ("review_tox_AMES", "AMES 예측 (0–1)")]:
                    target = next(i for i, row in enumerate(METRIC_GUIDE, 7) if row[0] == field)
                    self.assertIn(guide.cell(target, 3).value, columns[label].value)
                detail = {r[0].value: r[2].value for r in book["C000001"]}
                self.assertIn("↑ 높을수록", detail["평가 점수 (0–100점)"])
                self.assertIn("↓ 낮을수록", detail["AMES 예측 (0–1)"])
            finally:
                book.close()


if __name__ == "__main__":
    unittest.main()
